"""Read-only offline regression entry for the task and its saved report."""
import argparse
import hashlib
import importlib.util
import json
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True

"""Offline regression checks for the required workbook interface and penalties."""
import argparse
import copy
import importlib.util
import json
import tempfile
from pathlib import Path
from openpyxl import load_workbook
scoring_TASK = Path(__file__).resolve().parents[1]

def scoring_load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module
scoring_ENGINE = scoring_load_module('scoring_v2', scoring_TASK / 'tests/score_output.py')
scoring_FINALIZER = scoring_load_module('fixed_finalizer', scoring_TASK / 'tests/finalize.py')
scoring_KEY = json.loads((scoring_TASK / 'tests/results_key.json').read_text())
scoring_CRITERIA = scoring_KEY['rows'] + scoring_KEY['tables'] + scoring_KEY['checks']

def scoring_program_details(result):
    return {'kind': 'programmatic', 'score': round(result['score'], 6), 'criteria': [{'name': criterion['criterion_id'], 'description': criterion.get('description', criterion.get('label')), 'weight': abs(criterion['weight']), 'negate': criterion['weight'] < 0, 'raw': result['verdicts'][criterion['criterion_id']]['credit'], 'value': round(1 - result['verdicts'][criterion['criterion_id']]['credit'] if criterion['weight'] < 0 else result['verdicts'][criterion['criterion_id']]['credit'], 4), 'evidence': result['verdicts'][criterion['criterion_id']].get('evidence'), 'gate_reason': result['verdicts'][criterion['criterion_id']].get('gate_reason')} for criterion in scoring_CRITERIA]}

def scoring_evaluate(path):
    result = scoring_ENGINE.score(scoring_ENGINE.extract(path, scoring_KEY), scoring_KEY)
    pooled, counted, broken = scoring_FINALIZER.pooled_score({'reward': [scoring_program_details(result)]})
    assert not broken and counted == len(scoring_CRITERIA)
    assert abs(scoring_FINALIZER.apply_score_rules(pooled, result) - result['score']) < 1e-09
    return result

def scoring_mutation_cases():
    return [('both_sheets_renamed', 'Results、Detail 均改错表名', 0.882353, 'gate', None), ('results_columns_reversed', 'Results 四列逆序', 0.941176, 'cap', None), ('results_rows_reversed', 'Results 三行逆序', 0.941176, 'S01', None), ('units_eur', 'Results USD 改为 EUR', 0.941176, 'cap', None), ('detail_columns_reversed', 'Detail 三列逆序', 1.0, 'cap', None), ('sources_empty', '清空全部 Results source_note', 1.0, 'cap', None), ('sources_header_empty', '来源列连表头一并清空', 0.941176, 'cap', None), ('results_plan_d', 'Results 尾部新增 Plan D=999999', 1.0, 'cap', None), ('results_plus_cent', 'Results 三项各增加 USD 0.01', 1.0, 'numeric', None), ('results_missing', '删除必需 Results 表', None, 'gate', None), ('detail_missing', '删除必需 Detail 表', None, 'gate', None), ('results_header_case', 'Results 表头 value 改为 Value', None, 'cap', None), ('detail_header_case', 'Detail 表头 date 改为 Date', None, 'cap', None), ('results_extra_column', 'Results 增加第五列', None, 'positive', None), ('detail_extra_column', 'Detail 增加第四列', None, 'positive', None), ('detail_plan_d', 'Detail 尾部新增 Plan D', None, 'cap', None), ('results_duplicate', 'Results 重复 Plan A 记录', None, 'S05', None), ('detail_duplicate', 'Detail 重复 A 记录', None, 'P02', None), ('results_blank_label', 'Results 增加无标签金额记录', None, 'cap', None), ('detail_blank_label', 'Detail 增加无标签金额记录', None, 'cap', None), ('source_one_empty', '只清空 Results D3', None, 'cap', None), ('sources_numeric', '来源说明全部替换为金额数字', None, 'cap', None), ('sources_whitespace', '来源说明全部替换为空白', None, 'cap', None), ('results_general', 'Results 整数金额使用 General，存值已等于两位舍入', None, 'positive', None), ('detail_general', 'Detail 整数金额使用 General，存值已等于两位舍入', None, 'positive', None), ('money_one_decimal', '六个整数金额显示一位，存值已等于两位舍入', None, 'positive', None), ('money_three_decimals', '六个整数金额显示三位，存值已等于两位舍入', None, 'positive', None), ('detail_plus_cent', 'Detail 三项各增加 USD 0.01', None, 'numeric', None), ('units_case', 'Results USD 改为 usd', None, 'cap', None), ('results_blank_row', 'Results 首条结果前插入空行', None, 'S01', None), ('wrong_date', 'Detail A 日期改为 2008-01-01', None, 'V01', None), ('support_error', '支持表出现未计分的 Excel 错误', None, 'P01', None), ('header_trailing_spaces', '两个表的必需表头带前后空格', None, 'positive', None), ('extra_styled_column', '两个表增加空白但有样式的列', None, 'positive', None), ('additional_support_sheet', '增加自由组织支持表', None, 'positive', None), ('equivalent_storage', '六个存值为整数，数字格式为 General', None, 'positive', None), ('formatted_subcent', '存值多 0.001 USD，格式显示两位', None, 'positive', None), ('stored_cent', '存值多 0.004 USD，General 格式未显示两位', None, 'S06', None), ('three_decimal_subcent', '存值多 0.001 USD，格式显示三位', None, 'S06', None), ('sources_numeric_text', '来源说明替换为数值字符串', None, 'cap', None), ('results_sheet_case', '必需表名 Results 改为 results', None, 'gate', None), ('detail_sheet_case', '必需表名 Detail 改为 detail', None, 'gate', None), ('detail_plan_d_zero', '额外 Plan D 金额为零仍属业务记录', None, 'cap', None), ('results_unknown_label', 'Results 非规定标签保留业务金额', None, 'cap', None), ('results_label_case', 'Results 标签大小写不符但保留金额', None, 'cap', None)]

def scoring_mutate(book, name):
    results, detail = (book['Results'], book['Detail'])
    if name == 'both_sheets_renamed':
        results.title, detail.title = ('WrongResults', 'WrongDetail')
    elif name.endswith('columns_reversed'):
        sheet = results if name.startswith('results') else detail
        rows = [[copy.copy(cell) for cell in row] for row in sheet]
        for row_index, row in enumerate(rows, 1):
            for column_index, original in enumerate(reversed(row), 1):
                cell = sheet.cell(row_index, column_index, original.value)
                cell._style = copy.copy(original._style)
    elif name == 'results_rows_reversed':
        rows = list(results.values)[1:]
        for row_index, values in enumerate(reversed(rows), 2):
            for column_index, value in enumerate(values, 1):
                results.cell(row_index, column_index, value)
    elif name in {'units_eur', 'units_case'}:
        for row in range(2, 5):
            results.cell(row, 3, 'EUR' if name == 'units_eur' else 'usd')
    elif name in {'sources_empty', 'sources_header_empty', 'source_one_empty', 'sources_numeric', 'sources_whitespace'}:
        for row in [3] if name == 'source_one_empty' else range(2, 5):
            results.cell(row, 4).value = 42200 if name == 'sources_numeric' else '  ' if name == 'sources_whitespace' else None
        if name == 'sources_header_empty':
            results['D1'] = None
    elif name in {'results_missing', 'detail_missing'}:
        book.remove(results if name.startswith('results') else detail)
    elif name in {'results_header_case', 'detail_header_case'}:
        (results if name.startswith('results') else detail)['B1'] = 'Value' if name.startswith('results') else 'Date'
    elif name.endswith('extra_column'):
        sheet = results if name.startswith('results') else detail
        sheet.cell(1, sheet.max_column + 1, 'extra')
    elif name.endswith(('plan_d', 'duplicate', 'blank_label')):
        sheet = results if name.startswith('results') else detail
        label = 'Plan D principal balance' if sheet is results else 'D'
        values = [label, 999999, 'USD', 'Records.pdf'] if sheet is results else [label, '2007-12-31', 999999]
        if name.endswith('duplicate'):
            values = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
        if name.endswith('blank_label'):
            values[0] = None
        sheet.append(values)
        sheet.cell(sheet.max_row, 2 if sheet is results else 3).number_format = '#,##0.00'
    elif name.endswith('plus_cent'):
        sheet, column = (results, 2) if name.startswith('results') else (detail, 3)
        for row in range(2, 5):
            sheet.cell(row, column).value += 0.01
    elif name in {'results_general', 'detail_general', 'money_one_decimal', 'money_three_decimals'}:
        targets = [(results, 2), (detail, 3)]
        if name.endswith('general'):
            targets = [targets[0] if name.startswith('results') else targets[1]]
        number_format = 'General' if name.endswith('general') else '0.0' if name == 'money_one_decimal' else '0.000'
        for sheet, column in targets:
            for row in range(2, 5):
                sheet.cell(row, column).number_format = number_format
    elif name == 'results_blank_row':
        results.insert_rows(2)
    elif name == 'wrong_date':
        detail['B2'] = '2008-01-01'
    elif name == 'support_error':
        book.create_sheet('Working')['A1'] = '=1/0'
        book['Working']['A1'] = '#DIV/0!'
    elif name == 'header_trailing_spaces':
        for sheet in (results, detail):
            for cell in sheet[1]:
                cell.value = '  ' + cell.value + '  '
    elif name == 'extra_styled_column':
        for sheet in (results, detail):
            sheet.cell(1, sheet.max_column + 1)._style = copy.copy(sheet['A1']._style)
    elif name == 'additional_support_sheet':
        book.create_sheet('Support')['A1'] = '自由组织的支持工作表'
    elif name == 'equivalent_storage':
        for sheet, column in [(results, 2), (detail, 3)]:
            for row in range(2, 5):
                sheet.cell(row, column).value = int(sheet.cell(row, column).value)
                sheet.cell(row, column).number_format = 'General'
    elif name in {'formatted_subcent', 'stored_cent', 'three_decimal_subcent'}:
        results['B2'].value += 0.004 if name == 'stored_cent' else 0.001
        results['B2'].number_format = {'formatted_subcent': '0.00', 'stored_cent': 'General', 'three_decimal_subcent': '0.000'}[name]
    elif name == 'sources_numeric_text':
        results['D2'] = '42,200.00'
    elif name in {'results_sheet_case', 'detail_sheet_case'}:
        (results if name.startswith('results') else detail).title = 'results' if name.startswith('results') else 'detail'
    elif name == 'detail_plan_d_zero':
        detail.append(['D', '2007-12-31', 0])
    elif name in {'results_unknown_label', 'results_label_case'}:
        results['A2'] = 'Other principal balance' if name == 'results_unknown_label' else 'plan a principal balance'
    else:
        raise ValueError(name)

def scoring_run(write_cases):
    golden = scoring_TASK / 'solution/golden/answer.xlsx'
    assert scoring_evaluate(golden)['score'] == 1.0
    folder = scoring_TASK / 'validation/negative_cases'
    folder.mkdir(exist_ok=True)
    cases = []
    for name, mutation, old_score, expected, unused in scoring_mutation_cases():
        path = scoring_TASK / 'validation' / ('positive_cases' if expected == 'positive' else 'negative_cases') / (name + '.xlsx')
        if write_cases:
            book = load_workbook(golden)
            scoring_mutate(book, name)
            book.save(path)
            book.close()
        result = scoring_evaluate(path)
        penalties = [identifier for identifier, verdict in result['verdicts'].items() if verdict['weight'] < 0 and verdict['credit']]
        if expected == 'positive':
            assert result['score'] == 1.0, (name, result)
        else:
            assert result['score'] < 1.0, name
        if expected == 'cap':
            assert result['score'] <= 0.69 and result['interface_gate_passed'], name
            assert any((rule['triggered'] and rule['maximum_score'] == 0.69 for rule in result['score_rules'])), name
        elif expected == 'gate':
            assert result['score'] == 0 and (not result['interface_gate_passed']), name
        elif expected == 'numeric':
            prefix = 'R' if name.startswith('results') else 'D'
            assert all((not result['verdicts'][prefix + f'{index:02d}']['credit'] for index in range(1, 4))), name
        elif expected != 'positive':
            assert expected in penalties, (name, penalties)
        cases.append({'file': path.name, 'mutation': mutation, 'expected_outcome': expected, 'old_score_from_review': old_score, 'old_score_scope': 'reviewer programmatic subset' if old_score is not None else 'not reviewed', 'new_program_score': round(result['score'], 6), 'penalties_triggered': penalties, 'interface_gate_passed': result['interface_gate_passed'], 'interface_errors': result['interface_errors'], 'score_rules': result['score_rules'], 'category': 'positive_control' if expected == 'positive' else 'negative'})
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as temporary:
            graded = Path(temporary) / 'reward.json'
            graded.write_text(json.dumps({'programmatic': 1.0}))
            workbook = load_workbook(path, read_only=True)
            candidate_cell = workbook.sheetnames[0] + '!A1'
            workbook.close()
            details = {'reward': [scoring_program_details(result), {'kind': 'agent', 'criteria': [{'name': identifier, 'weight': 3, 'value': 1.0, 'negate': True, 'reasoning': candidate_cell + ' is the candidate header used by this program-policy control.'} for identifier in ('L01', 'L02', 'L03')]}]}
            graded.with_name('reward-details.json').write_text(json.dumps(details))
            finalized = scoring_FINALIZER.compute(SimpleNamespace(graded=str(graded), graded_rc=0, candidate=str(path)))
            assert finalized['reward'] == round(result['score'], 6), name
            assert finalized['criteria_counted'] == 20 and (not finalized['verifier_error']), name
            recorded = json.loads(graded.with_name('reward-details.json').read_text())
            assert recorded['score_rules'] == result['score_rules'], name
            assert recorded['pooling']['before_rules'] == finalized['score_before_rules'], name
    with tempfile.TemporaryDirectory() as temporary:
        path = Path(temporary) / 'control.xlsx'
        book = load_workbook(golden)
        book.create_sheet('Support')['A1'] = '自由组织的支持工作表'
        book.save(path)
        assert scoring_evaluate(path)['score'] == 1.0
        for difference, expected in [(0.0049, 1), (0.005, 1), (0.0051, 0)]:
            book = load_workbook(golden)
            book['Results']['B2'].value += difference
            book.save(path)
            assert scoring_evaluate(path)['verdicts']['R01']['credit'] == expected
        assert scoring_evaluate(Path(temporary) / 'missing.xlsx')['score'] == 0
    engine = scoring_ENGINE.base.engine
    assert engine.precision_matches({'raw': 1.2345, 'number_format': 'General'}, 3, at_least=True)
    assert engine.precision_matches({'raw': 1e-05, 'number_format': 'General'}, 3, at_least=True)
    assert not engine.precision_matches({'raw': 1.2345, 'number_format': 'General'}, 3)
    assert engine.precision_matches({'raw': 1.2345, 'number_format': '0.000'}, 3)
    if write_cases:
        (folder / 'cases.json').write_text(json.dumps({'baseline_program_score': 1.0, 'cases': cases, 'precision_rule': 'p decimals: format shows p decimals OR stored value equals round(value,p); at least p also accepts stored precision >=p.', 'controls': ['additional_support_sheet=1', 'tolerance_0.0049=accepted', 'tolerance_0.005=accepted', 'tolerance_0.0051=rejected', 'missing_output=0'], 'llm_status': 'NOT_RUN'}, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'negative_cases_passed': sum((case['category'] == 'negative' for case in cases)), 'positive_controls_passed': sum((case['category'] == 'positive_control' for case in cases)), 'finalize_cases_passed': len(cases), 'numeric_boundary_controls': 5, 'golden': 1.0}))

"""Regression checks for candidate-access evidence and saved judgments."""
import argparse
import copy
import importlib.util
import json
import subprocess
import sys
import tempfile
from pathlib import Path
access_TASK = Path(__file__).resolve().parents[1]

def access_load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def access_main():
    finalizer = access_load('artifact_finalizer', access_TASK / 'tests/finalize.py')
    checker = access_load('artifact_checker', access_TASK / 'tests/llm_artifact_check.py')
    golden = access_TASK / 'solution/golden_output/answer.xlsx'
    blind = json.loads((access_TASK / 'history/validation/judge_fixtures/saved_blind.json').read_text())
    with tempfile.TemporaryDirectory() as temporary:
        folder = Path(temporary)
        graded = folder / 'graded/reward.json'
        graded.parent.mkdir()
        graded.write_text('{"programmatic": 1.0}')
        valid = {'reward': [{'kind': 'programmatic', 'criteria': [{'name': 'control', 'weight': 10, 'value': 1.0, 'negate': False}]}, {'kind': 'agent', 'criteria': [{'id': name, 'weight': 3, 'value': 1.0, 'negate': True, 'reasoning': 'Results!D2 contains a source citation with no unsupported assertion.'} for name in ['L01', 'L02', 'L03']]}]}
        cases = [('evidence_backed_control', valid, False), ('saved_blind_C_golden_r4', {'reward': [valid['reward'][0], blind['agent']]}, True)]
        for name, reason, expected in [('no_cell', 'No unsupported assertions were found.', True), ('source_only_cell', 'The source Text!A131 supports this conclusion.', True), ('nonexistent_candidate_cell', 'Results!ZZ9999 supports this conclusion.', True), ('missing_with_cell', 'Results!D2 cannot be read because candidate workbook is missing.', True), ('unreadable_with_cell', 'Results!D2：无法解析候选文件。', True), ('missing_source_not_artifact', '候选 Results!D2 引用了不存在的来源文件。', False), ('missing_source_english', 'The candidate workbook Results!D2 cites a missing source file.', False), ('quoted_cell', "'Results'!$D$2 contains no audit-completion claim.", False)]:
            detail = copy.deepcopy(valid)
            detail['reward'][1]['criteria'][0]['reasoning'] = reason
            cases.append((name, detail, expected))
        for name, detail, expected_error in cases:
            graded.with_name('reward-details.json').write_text(json.dumps(detail))
            output = folder / name / 'reward.json'
            subprocess.run([sys.executable, str(access_TASK / 'tests/finalize.py'), '--graded', str(graded), '--graded-rc', '0', '--candidate', str(golden), '--out', str(output)], check=True)
            receipt = json.loads(output.read_text())
            assert bool(receipt['verifier_error']) == expected_error, (name, receipt)
            if expected_error:
                assert receipt['reward'] == 0
                exit_message = json.loads(output.with_name('reward_exit_message.json').read_text())
                assert exit_message['exit_code'] == 'judge:artifact_access_error', name
            else:
                assert receipt['reward'] == 1.0, name
                assert not output.with_name('reward_exit_message.json').exists(), name
        saved_valid = 0
        for fixture in (access_TASK / 'history/validation/judge_fixtures').glob('saved_valid_*.json'):
            payload = json.loads(fixture.read_text())
            assert checker.evidence_errors({'reward': payload['agent']}, golden) == [], fixture.name
            saved_valid += 1
        assert checker.evidence_errors({'reward': blind['agent']}, folder / 'missing.xlsx') == []
    print(json.dumps({'artifact_access_regressions': len(cases), 'saved_blind_cases': 1, 'saved_valid_cases': saved_valid, 'missing_candidate_keeps_program_gate': True}))


def check_saved_cases(task):
    engine = import_module('case_inventory_engine', task / 'tests/score_output.py')
    key = json.loads((task / 'tests/results_key.json').read_text())
    counts = {}
    for category in ('negative_cases', 'positive_cases'):
        folder = task / 'validation' / category
        cases = json.loads((folder / 'cases.json').read_text())['cases']
        assert {entry['file'] for entry in cases} == {path.name for path in folder.glob('*.xlsx')}
        for entry in cases:
            result = engine.score(engine.extract(folder / entry['file'], key), key)
            expected = entry.get('new_program_score', entry.get('program_score', entry.get('expected_program_score')))
            assert expected is not None and abs(result['score'] - expected) < 0.000001, entry['file']
            assert result['score'] == 1.0 if category == 'positive_cases' else result['score'] < 1.0
        counts[category] = len(cases)
    return counts


def check_saved_report(task, track):
    engine = import_module('saved_report_engine', task / 'tests/score_output.py')
    key = json.loads((task / 'tests/results_key.json').read_text())
    golden = track / 'golden/candidate.xlsx'
    assert golden.read_bytes() == (task / 'solution/golden_output/answer.xlsx').read_bytes()
    assert engine.score(engine.extract(golden, key), key)['score'] == 1.0
    reward = json.loads((track / 'golden/reward.json').read_text())
    judge = json.loads((track / 'golden/judge.json').read_text())
    consistency = json.loads((track / 'consistency.json').read_text())
    assert reward['verifier_error'] == 1 and reward['score_available'] is False
    assert judge['score_final'] is None and judge['judge_completed'] is False
    assert consistency['status'] == 'NO_VALID_NEW_JUDGMENT'
    assert all(group['n'] == 0 and not group['run_ids'] and not group['requirement_met']
               for group in consistency['groups'])
    summary = json.loads((track / 'history/golden_judge_fix_summary.json').read_text())
    archive = track / 'history' / summary['history']
    assert (archive / 'golden/reward.json').is_file()
    assert len(list((archive / 'judge_runs').glob('golden_r[1-5]/reward-details.json'))) == 5
    for folder in sorted((track / 'rollouts').iterdir()):
        candidate = folder / 'candidate.xlsx'
        result = engine.score(engine.extract(candidate, key), key)
        receipt = json.loads((folder / 'reward.json').read_text())
        saved_program = receipt.get('programmatic_score', receipt['reward'])
        assert abs(result['score'] - saved_program) < 0.000001, folder.name
        assert receipt['reward'] == summary['new_model_scores'][folder.name], folder.name
    print(json.dumps({'saved_report_checked': track.parent.name, 'programmatic_golden': 1.0,
                      'valid_new_judgments': 0, 'current_golden_score': None,
                      'archived_judgments_excluded': 5}))


def import_module(name, path):
    specification = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(module)
    return module


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--trajectory', type=Path)
    arguments = parser.parse_args()
    task = Path(__file__).resolve().parents[1]
    track = arguments.trajectory or task.parents[1] / 'ZCloud-市政财务-002-20261022-轨迹报告/muskegon-acceleration-clause'
    counts = check_saved_cases(task)
    with tempfile.TemporaryDirectory() as temporary:
        scoring_run(False)
        access_main()
    check_saved_report(task, track)
    print(json.dumps(dict(status='PASS', **counts)))


if __name__ == '__main__':
    main()
