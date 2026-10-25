"""Render each reference sheet using Calc PDF export and pdftoppm."""
import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def render(golden_folder):
    source = golden_folder / 'candidate.xlsx'
    renders = golden_folder / 'renders'
    renders.mkdir(exist_ok=True)
    outputs = []
    with tempfile.TemporaryDirectory(dir=golden_folder) as temporary:
        temporary = Path(temporary)
        for index, name in enumerate(load_workbook(source).sheetnames, 1):
            book = load_workbook(source)
            for sheet in list(book):
                if sheet.title != name:
                    book.remove(sheet)
            sheet = book[name]
            sheet.print_area = sheet.dimensions
            sheet.page_setup.orientation = 'landscape'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_margins.left = sheet.page_margins.right = 0.35
            sheet.page_margins.top = sheet.page_margins.bottom = 0.35
            stem = f'sheet-{index:02d}'
            workbook = temporary / (stem + '.xlsx')
            book.save(workbook)
            book.close()
            subprocess.run(['soffice', '-env:UserInstallation=' + (temporary / 'profile').resolve().as_uri(),
                            '--headless', '--convert-to', 'pdf:calc_pdf_Export',
                            '--outdir', str(temporary), str(workbook)],
                           check=True, capture_output=True, timeout=60)
            pdf = temporary / (stem + '.pdf')
            assert pdf.is_file(), stem
            shutil.copyfile(pdf, renders / pdf.name)
            subprocess.run(['pdftoppm', '-png', '-singlefile', '-r', '110',
                            str(renders / pdf.name), str(renders / stem)],
                           check=True, capture_output=True, timeout=30)
            outputs.append({'sheet': name, 'pdf': f'renders/{stem}.pdf', 'png': f'renders/{stem}.png'})
    result = {'format': 'xlsx', 'reviewed_at': datetime.now().astimezone().isoformat(timespec='seconds'),
              'renderer': 'LibreOffice calc_pdf_Export + pdftoppm',
              'candidate_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
              'scope': '逐表从更新后的参考工作簿导出；临时副本使用 A3 横向和单页缩放。',
              'worksheet_count': len(outputs), 'render_count': len(outputs), 'renders': outputs,
              'render_status': 'COMPLETED', 'visual_inspection_status': 'PENDING',
              'visual_findings': [], 'source_workbook_unchanged': True,
              'scoring_revision': '评分器 v2（2026-10-26）'}
    (golden_folder / 'visual_review.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'rendered_sheets': len(outputs), 'renderer': result['renderer']}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('golden_folder', type=Path)
    render(parser.parse_args().golden_folder)
