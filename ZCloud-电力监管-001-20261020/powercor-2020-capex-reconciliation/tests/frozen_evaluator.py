import argparse, json, math, re, sys, unicodedata
from pathlib import Path
HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
EXCEL_ERRORS = {'#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A', '#NULL!', '#NUM!', '#SPILL!', '#CALC!', '#GETTING_DATA'}
YES = {'yes', 'y', 'true', '1', 't'}
NO = {'no', 'n', 'false', '0', 'f'}

def sha256_file(path) -> str:
    import hashlib
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()

def is_error_guard() -> bool:
    return all((not is_excel_error(x) for x in ('#2A', '#13', '#1 base bid'))) and is_excel_error('#REF!')

def is_excel_error(v) -> bool:
    return isinstance(v, str) and v.strip().upper() in EXCEL_ERRORS

def norm_label(s) -> str:
    if s is None:
        return ''
    t = unicodedata.normalize('NFKC', str(s)).casefold()
    t = t.replace('’', "'").replace('‘', "'").replace('“', '"').replace('”', '"')
    t = re.sub('[‐-―−]', '-', t)
    t = re.sub('\\s+', ' ', t).strip()
    return t.rstrip(':.;').strip()

def parse_number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if not isinstance(v, str):
        return None
    t = unicodedata.normalize('NFKC', v).strip()
    if not t or t.upper() in EXCEL_ERRORS:
        return None
    neg = t.startswith('(') and t.endswith(')')
    t = t.strip('()')
    pct = '%' in t
    t = re.sub('(?i)\\b(usd|gbp|eur|us\\$)\\b', '', t)
    t = re.sub('[,$£€%\\s]', '', t)
    m = re.fullmatch('[-+]?\\d*\\.?\\d+(?:[eE][-+]?\\d+)?', t)
    if not m:
        return None
    x = float(t)
    if not math.isfinite(x):
        return None
    if neg:
        x = -x
    return ('pct', x) if pct else x

def norm_choice(v) -> str:
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'yes' if v else 'no'
    if isinstance(v, (int, float)) and float(v).is_integer():
        v = int(v)
    t = unicodedata.normalize('NFKC', str(v)).casefold().strip().rstrip('.')
    return re.sub('[\\s\\-_]+', '_', t)

def number_matches(cand, gold, tol, percent_forms=False) -> bool:
    x = parse_number(cand)
    if x is None:
        return False
    if isinstance(x, tuple):
        if not percent_forms:
            return False
        x = x[1]
    a, r = (float(tol.get('abs', 0) or 0), float(tol.get('rel', 0) or 0))
    return abs(x - gold) <= max(a, r * abs(gold), 1e-09)

def value_matches(row, cand, gold) -> bool:
    kind = row.get('type', 'number')
    if cand is None or (isinstance(cand, str) and (not cand.strip() or cand.strip().upper() in EXCEL_ERRORS)):
        return False
    if kind in ('number', 'integer'):
        tol = row.get('tolerance') or ({'abs': 0.5} if kind == 'integer' else {'abs': 0.005, 'rel': 1e-06})
        percent = bool(row.get('percent_forms'))
        if isinstance(parse_number(cand), tuple) and norm_label(row.get('unit')) in {'%', 'percent', 'percentage', 'percentage points'}:
            percent = True
        return number_matches(cand, float(gold), tol, percent)
    if kind == 'yesno':
        c = norm_choice(cand)
        g = norm_choice(gold)
        return c in YES and g in YES or (c in NO and g in NO)
    if kind == 'choice':
        return norm_choice(cand) == norm_choice(gold)
    if kind == 'text_set':
        split = lambda s: {norm_choice(p) for p in re.split('[;,\\n]+', str(s)) if p.strip()}
        return split(cand) == ({norm_choice(g) for g in gold} if isinstance(gold, list) else split(gold))
    raise ValueError(f'unknown row type {kind!r}')

def displayed_percent(value, number_format):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return value
    clean = re.sub('"[^"]*"|\\\\.|_.|\\*.', '', number_format or '')
    sections = clean.split(';')[:3]
    conditions = [re.search('\\[(<=|>=|<>|=|<|>)(-?\\d+(?:\\.\\d+)?)\\]', s) for s in sections]
    selected = None
    if any(conditions):
        comparisons = {'<': lambda a, b: a < b, '>': lambda a, b: a > b, '<=': lambda a, b: a <= b, '>=': lambda a, b: a >= b, '=': lambda a, b: a == b, '<>': lambda a, b: a != b}
        for section, condition in zip(sections, conditions):
            if condition is None or comparisons[condition[1]](value, float(condition[2])):
                selected = section
                break
    else:
        idx = 1 if value < 0 and len(sections) > 1 else 2 if value == 0 and len(sections) > 2 else 0
        selected = sections[idx] if sections else ''
    scales = re.sub('\\[[^\\]]*\\]', '', selected or '').count('%')
    return f'{value * 100 ** scales:.17g}%' if scales else value

def read_static(path: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = [[displayed_percent(c.value, c.number_format) for c in r] for r in ws.iter_rows()]
    wb.close()
    return out

def read_formulas(path: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = [[c if isinstance(c, str) and c.startswith('=') else None for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out

def find_sheet(book: dict, name: str):
    for title, rows in book.items():
        if norm_label(title) == norm_label(name):
            return rows
    return None

def header_map(rows, wanted):
    if not rows:
        return None
    hdr = {norm_label(v): i for i, v in enumerate(rows[0]) if v is not None}
    return {w: hdr[norm_label(w)] for w in wanted if norm_label(w) in hdr}

def result_label(value, wanted):
    label = norm_label(value)
    bare = norm_label(re.sub('\\s*\\[[^\\]]*\\]\\s*$', '', str(value)))
    return bare if label not in wanted and bare in wanted else label

def result_columns(key, rows):
    if not rows:
        return None
    cols = header_map(rows, ['label', 'value'])
    if len(cols) == 2:
        return cols
    alias = header_map(rows, ['item', 'value'])
    if len(alias) == 2:
        return {'label': alias['item'], 'value': alias['value']}
    wanted = {norm_label(r['label']) for r in key['rows']}
    counts = [(sum((result_label(r[ci], wanted) in wanted for r in rows[1:] if ci < len(r))), ci) for ci in range(max(map(len, rows), default=0))]
    counts.sort(reverse=True)
    if counts and counts[0][0] >= max(1, len(wanted) // 2):
        if len(counts) == 1 or counts[0][0] > counts[1][0]:
            return {'label': counts[0][1], 'value': counts[0][1] + 1}
    return None

def score_rows(key: dict, book: dict, formulas: dict | None=None) -> dict:
    verdicts = {}
    rows = find_sheet(book, key.get('sheet', 'Results'))
    cols = result_columns(key, rows)
    frows = find_sheet(formulas or {}, key.get('sheet', 'Results'))
    written, was_formula, conflicts = ({}, {}, set())
    if rows and cols and (len(cols) == 2):
        for i, r in enumerate(rows[1:], 1):
            if len(r) > max(cols.values()) and r[cols['label']] is not None:
                wanted = {norm_label(x['label']) for x in key['rows']}
                lab = result_label(r[cols['label']], wanted)
                if lab in written and written[lab] != r[cols['value']]:
                    conflicts.add(lab)
                if lab not in written:
                    written[lab] = r[cols['value']]
                    fr = frows[i] if frows and i < len(frows) else None
                    was_formula[lab] = bool(fr and len(fr) > cols['value'] and fr[cols['value']])
    base = {}
    for row in key['rows']:
        cid = row['criterion_id']
        if rows is None or not cols or len(cols) < 2:
            base[cid] = ('not_met', 'the Results sheet or its label/value header is absent', None)
            continue
        lab = norm_label(row['label'])
        if lab not in written:
            base[cid] = ('not_met', 'no row carries this label', None)
            continue
        if lab in conflicts:
            base[cid] = ('unresolved', 'conflicting duplicate labels require adjudication', None)
            continue
        cand = written[lab]
        golds = [row['gold']] + list(row.get('accepted') or [])
        ok = any((value_matches(row, cand, g) for g in golds))
        live = was_formula.get(lab, False)
        base[cid] = ('met' if ok else 'not_met', f"written {cand!r} as {('a formula' if live else 'a literal')}", cand)
    chosen = {}
    by_id = {r['criterion_id']: r for r in key['rows']}
    for alt in key.get('alternative_sets') or []:
        ids = [i for i in alt['gold'] if i in by_id]
        w = lambda st: sum((by_id[i]['weight'] for i in ids if st[i] == 'met'))
        alt_status = {i: 'met' if base[i][2] is not None and value_matches(by_id[i], base[i][2], alt['gold'][i]) else 'not_met' for i in ids}
        if w(alt_status) > w({i: base[i][0] for i in ids}):
            for i in ids:
                base[i] = (alt_status[i], base[i][1] + f" (judged under '{alt['name']}')", base[i][2])
            chosen[alt['name']] = ids
    for cid, (st, why, _) in base.items():
        verdicts[cid] = {'status': st, 'decided_by': 'program', 'evidence': why}
    return (verdicts, chosen)

def table_records(rows, keycols, valcols, expected_keys):
    cols = header_map(rows, keycols + valcols) if rows else None
    if not rows or not cols or any((k not in cols for k in keycols + valcols)):
        return (None, {})

    def same(w, g):
        return w == g or (w.startswith(g) and (not w[len(g):len(g) + 1].isalnum()))
    grouped = {}
    for row in rows[1:]:
        row = list(row) + [None] * (max(cols.values()) + 1 - len(row))
        if any((row[cols[k]] is not None for k in keycols)):
            written = tuple((norm_choice(row[cols[k]]) for k in keycols))
            grouped.setdefault(written, []).append(row)
    for written in list(grouped):
        if written not in expected_keys:
            fits = [g for g in expected_keys if len(g) == len(written) and all((same(w, v) for w, v in zip(written, g)))]
            if len(fits) == 1 and fits[0] not in grouped:
                grouped[fits[0]] = grouped.pop(written)
    return (cols, grouped)

def score_tables(key: dict, book: dict) -> dict:
    verdicts = {}
    for t in key.get('tables') or []:
        cid = t['criterion_id']
        rows = find_sheet(book, t['sheet'])
        keycols, valcols = (t['key'], list((t.get('value_columns') or {}).keys()))
        gold = {tuple((norm_choice(g[k]) for k in keycols)): g for g in t['gold']}
        cols, got = table_records(rows, keycols, valcols, gold)
        if (t.get('pass_rule') or {}).get('per_record') is True:
            verdicts[cid] = score_table_per_record(t, cols, got, gold)
            continue
        if cols is None:
            verdicts[cid] = {'status': 'not_met', 'decided_by': 'program', 'evidence': f"sheet {t['sheet']!r} or its contracted header is absent"}
            continue
        hit = [k for k in gold if k in got]
        extra = [k for k in got if k not in gold]
        duplicates = ['/'.join(k) for k, records in got.items() if len(records) > 1]
        wrong, ambiguous = ([], [])
        for k in hit:
            for vc, tol in (t.get('value_columns') or {}).items():
                spec = {'type': tol.get('type', 'number'), 'tolerance': tol, 'percent_forms': tol.get('percent_forms'), 'unit': (t.get('units') or {}).get(vc)}
                outcomes = [value_matches(spec, record[cols[vc]], gold[k][vc]) for record in got[k]]
                field = f"{'/'.join(k)}:{vc}"
                if not any(outcomes):
                    wrong.append(field)
                elif not all(outcomes):
                    ambiguous.append(field)
        rule = t.get('pass_rule') or {}
        recall = len(hit) / len(gold) if gold else 1.0
        structural_ok = recall >= rule.get('min_recall', 1.0) and (rule.get('ignore_extra') is True or len(extra) <= rule.get('max_extra', 0)) and (len(duplicates) <= rule.get('max_duplicate_keys', math.inf))
        max_wrong = rule.get('max_wrong_values', 0)
        status = 'not_met' if not structural_ok or len(wrong) > max_wrong else 'met' if len(wrong) + len(ambiguous) <= max_wrong else 'unresolved'
        verdicts[cid] = {'status': status, 'decided_by': 'program', 'duplicate_keys': duplicates, 'ambiguous_values': ambiguous, 'evidence': f'{len(hit)}/{len(gold)} keyed rows found, {len(extra)} not in gold, wrong values: {wrong[:6]}; duplicate keys: {duplicates[:6]}; ambiguous duplicate fields: {ambiguous[:6]}'}
    return verdicts

def score_table_per_record(table: dict, cols, got: dict, gold: dict) -> dict:
    if not gold or len(gold) != len(table['gold']):
        raise ValueError('per_record requires nonempty, unique reference keys')
    errors = []
    for record_key, reference in gold.items():
        identifier = list(record_key)
        if cols is None or record_key not in got:
            errors.append({'key': identifier, 'reason': 'missing record'})
            continue
        wrong_fields = []
        for column, tolerance in (table.get('value_columns') or {}).items():
            spec = {'type': tolerance.get('type', 'number'), 'tolerance': tolerance, 'percent_forms': tolerance.get('percent_forms'), 'unit': (table.get('units') or {}).get(column)}
            if not all((value_matches(spec, record[cols[column]], reference[column]) for record in got[record_key])):
                wrong_fields.append(column)
        if wrong_fields:
            errors.append({'key': identifier, 'reason': 'incorrect or conflicting values', 'fields': wrong_fields})
    total = len(gold)
    correct = total - len(errors)
    return {'status': 'met' if correct == total else 'not_met', 'decided_by': 'program', 'credit_fraction': correct / total, 'correct_records': correct, 'total_records': total, 'error_records': errors, 'duplicate_keys': ['/'.join(key) for key, records in got.items() if len(records) > 1], 'evidence': f'{correct}/{total} records correct; error records: {errors}'}

def norm_unit(v) -> str:
    if v is None:
        return ''
    return re.sub('[^\\w%$£€/]+', '', unicodedata.normalize('NFKC', str(v)).casefold())

def _cell(row, idx):
    return row[idx] if row is not None and idx is not None and (idx < len(row)) else None

def _first_rows(key: dict, book: dict):
    rows = find_sheet(book, key.get('sheet', 'Results'))
    cols = result_columns(key, rows) if rows else None
    if not rows or not cols or len(cols) < 2:
        return (rows, None, {}, {})
    wanted = {norm_label(r['label']) for r in key['rows']}
    first = {}
    for r in rows[1:]:
        raw = _cell(r, cols['label'])
        if raw is not None:
            first.setdefault(result_label(raw, wanted), r)
    return (rows, cols, first, header_map(rows, ['unit', 'source_note']))

def _number(v):
    x = parse_number(v)
    return x[1] if isinstance(x, tuple) else x

def graded_error_free(key: dict, book: dict) -> set:
    graded = set()
    rows, cols, first, _ = _first_rows(key, book)
    if rows is not None and cols:
        title = next((t for t, v in book.items() if v is rows))
        for i, r in enumerate(rows):
            if i and _cell(r, cols['label']) is not None:
                graded.add((title, i, cols['value']))
    for t in key.get('tables') or []:
        sheet = find_sheet(book, t['sheet'])
        if sheet is None:
            continue
        title = next((n for n, v in book.items() if v is sheet))
        hdr = header_map(sheet, list((t.get('value_columns') or {}).keys()))
        for i in range(1, len(sheet)):
            for ci in hdr.values():
                graded.add((title, i, ci))
    return graded

def score_checks(key: dict, book: dict) -> dict:
    out = {}
    if not key.get('checks'):
        return out
    rows, cols, first, hmap = _first_rows(key, book)
    contract = [r for r in key['rows'] if r['weight'] >= 0]
    labels = [norm_label(r['label']) for r in contract]
    wanted = {norm_label(r['label']) for r in key['rows']}
    for c in key['checks']:
        kind, cid = (c['kind'], c['criterion_id'])
        ok, why = (False, '')
        if kind == 'results_header':
            want = [norm_label(x) for x in c.get('columns') or ['label', 'value', 'unit', 'source_note']]
            got = [norm_label(x) for x in (rows[0][:len(want)] if rows else [])]
            ok = got == want
            why = f'row 1 reads {got!r}, contract {want!r}'
        elif kind == 'results_rows_in_order':
            if cols:
                seq = [lab for r in rows[1:] for lab in [result_label(_cell(r, cols['label']), wanted)] if _cell(r, cols['label']) is not None and lab in wanted]
                ok = seq == labels
                why = f'{len(seq)} contract labels found in sheet order; contract has {len(labels)}'
            else:
                why = 'the Results sheet or its label/value header is absent'
        elif kind == 'results_units':
            if cols and hmap.get('unit') is not None:
                units = c.get('units') or {r['label']: r.get('unit') for r in contract}
                bad = [lab for lab, u in units.items() if norm_unit(_cell(first.get(norm_label(lab)), hmap['unit'])) != norm_unit(u)]
                ok = not bad
                why = f'{len(units) - len(bad)}/{len(units)} units as stated; wrong or missing: {bad[:5]}'
            else:
                why = 'no Results rows or no unit column'
        elif kind == 'results_source_notes':
            if cols and hmap.get('source_note') is not None:
                bad = [r['label'] for r in contract if not re.search('[0-9A-Za-z\\u00c0-\\uffff]', str(_cell(first.get(norm_label(r['label'])), hmap['source_note']) or ''))]
                ok = not bad
                why = f'{len(contract) - len(bad)}/{len(contract)} rows carry a source_note; empty: {bad[:5]}'
            else:
                why = 'no Results rows or no source_note column'
        elif kind == 'sum_relation':
            tol = float((c.get('tolerance') or {}).get('abs', 0) or 0)
            try:
                tv = _number(_cell(first.get(norm_label(c['target'])), cols['value'])) if cols else None
                parts = [(t['coef'], _number(_cell(first.get(norm_label(t['label'])), cols['value'])) if cols else None) for t in c['terms']]
            except (KeyError, TypeError):
                tv, parts = (None, [(0, None)])
            if tv is None or any((v is None for _, v in parts)):
                why = 'a term of the relation is missing or not numeric'
            else:
                resid = tv - sum((k * v for k, v in parts))
                ok = abs(resid) <= tol
                why = f"{c['target']} minus the weighted terms = {resid:.9g} (allowed {tol:g})"
        elif kind == 'ratio_relation':
            tol = float((c.get('tolerance') or {}).get('abs', 0) or 0)
            vals = [_number(_cell(first.get(norm_label(c[f])), cols['value'])) if cols else None for f in ('target', 'numerator', 'denominator')]
            if any((v is None for v in vals)) or vals[2] == 0:
                why = 'a term of the relation is missing, not numeric or zero'
            else:
                resid = vals[0] - float(c.get('scale', 1)) * vals[1] / vals[2]
                ok = abs(resid) <= tol
                why = f"{c['target']} minus scale x numerator / denominator = {resid:.9g} (allowed {tol:g})"
        elif kind == 'extra_numeric_results_rows':
            extra, absent = ([], [])
            if cols:
                for r in rows[1:]:
                    raw = _cell(r, cols['label'])
                    if raw is None or not str(raw).strip():
                        continue
                    if result_label(raw, wanted) not in wanted and _number(_cell(r, cols['value'])) is not None:
                        extra.append(str(raw)[:50])
                absent = [lab for lab in dict.fromkeys(labels) if lab not in first]
            ok = len(extra) > len(absent)
            why = f'numeric rows with labels outside the contract: {extra[:5]} ({len(extra)}); contract labels with no row: {len(absent)}; unrequested rows beyond them: {max(0, len(extra) - len(absent))}'
        elif kind == 'error_cells':
            graded = graded_error_free(key, book)
            bad = [f'{t}!R{i + 1}C{j + 1}' for t, grid in book.items() for i, r in enumerate(grid) for j, v in enumerate(r) if is_excel_error(v) and (t, i, j) not in graded]
            ok = bool(bad)
            why = f'error values outside the graded cells: {bad[:5]}'
        else:
            raise ValueError(f'unknown check kind {kind!r}')
        if c['weight'] < 0:
            out[cid] = {'status': 'triggered' if ok else 'clear', 'decided_by': 'program', 'evidence': why}
        else:
            out[cid] = {'status': 'met' if ok else 'not_met', 'decided_by': 'program', 'evidence': why}
    return out
CORE_DIMENSIONS = {'data_accuracy', 'professional_standard', 'factuality', 'content_quality'}

def criterion_meta(key: dict) -> dict:
    meta = {}
    for kind in ('rows', 'tables', 'ai_criteria', 'checks'):
        for c in key.get(kind) or []:
            meta[c['criterion_id']] = {'weight': c['weight'], 'dimension': c.get('dimension'), 'kind': kind}
    return meta

def aggregate(key: dict, verdicts: dict) -> dict:
    items = [(r['criterion_id'], r['weight'], bool(r.get('critical'))) for r in key['rows']] + [(t['criterion_id'], t['weight'], bool(t.get('critical'))) for t in key.get('tables') or []] + [(a['criterion_id'], a['weight'], bool(a.get('critical'))) for a in key.get('ai_criteria') or []] + [(c['criterion_id'], c['weight'], False) for c in key.get('checks') or []]
    penalties = [(c, w) for c, w, _ in items if w < 0]
    items = [i for i in items if i[1] >= 0]
    total = sum((w for _, w, _ in items))
    st = lambda c: verdicts.get(c, {}).get('status', 'pending')
    per_record = {table['criterion_id'] for table in key.get('tables') or [] if (table.get('pass_rule') or {}).get('per_record') is True}

    def credit(criterion_id):
        verdict = verdicts.get(criterion_id, {})
        if criterion_id in per_record and 'credit_fraction' in verdict:
            return verdict['credit_fraction']
        return 1 if st(criterion_id) == 'met' else 0

    def points(criterion_id, weight):
        if criterion_id in per_record:
            return weight * credit(criterion_id)
        return weight if st(criterion_id) == 'met' else 0
    earned = sum((points(c, w) for c, w, _ in items if st(c) == 'met' or c in per_record))
    open_w = sum((w for c, w, _ in items if st(c) in ('pending', 'unresolved')))
    crit_fail = [c for c, _, k in items if k and st(c) == 'not_met']
    crit_open = [c for c, _, k in items if k and st(c) in ('pending', 'unresolved')]
    critical = 'FAIL' if crit_fail else 'UNRESOLVED' if crit_open else 'PASS'
    lower, upper = (earned / total, (earned + open_w) / total)
    pen_hit, pen_w, pen_open = ([], 0, 0)
    if penalties:
        pen_hit = [c for c, w in penalties if st(c) in ('triggered', 'met')]
        pen_w = sum((-w for c, w in penalties if st(c) in ('triggered', 'met')))
        pen_open = sum((-w for c, w in penalties if st(c) in ('pending', 'unresolved')))
        lower = max(0.0, (earned - pen_w - pen_open) / total)
        upper = max(0.0, (earned + open_w - pen_w) / total)
    if crit_fail or upper < 0.7:
        accepted = 'FAIL'
    elif (open_w or pen_open) and (not (lower >= 0.7 and (not crit_open))):
        accepted = 'UNRESOLVED'
    else:
        accepted = 'PASS' if lower >= 0.7 and critical == 'PASS' else 'UNRESOLVED' if open_w or pen_open else 'FAIL'
    out = {'total_weight': total, 'earned': earned, 'open_weight': open_w, 'score': None if open_w or pen_open else lower, 'lower_bound': round(lower, 6), 'upper_bound': round(upper, 6), 'critical_status': critical, 'critical_failed': crit_fail, 'critical_unresolved': crit_open, 'accepted': accepted, 'failed': [c for c, _, _ in items if st(c) == 'not_met']}
    if penalties:
        out['penalty_weight'] = pen_w
        out['penalties_triggered'] = pen_hit
    meta = criterion_meta(key)
    if any((m['dimension'] for m in meta.values())):
        by = {}
        for cid, m in meta.items():
            if m['weight'] <= 0:
                continue
            b = by.setdefault(m['dimension'] or 'unspecified', {'weight': 0, 'earned': 0, 'open': 0})
            b['weight'] += m['weight']
            b['earned'] += points(cid, m['weight'])
            b['open'] += m['weight'] if st(cid) in ('pending', 'unresolved') else 0
        out['by_dimension'] = by
        core = [b for d, b in by.items() if d in CORE_DIMENSIONS]
        core_w, core_e, core_o = (sum((b[k] for b in core)) for k in ('weight', 'earned', 'open'))
        out['core_weight'] = core_w
        out['core_score'] = None if not core_w or core_o else core_e / core_w
    return out

def formula_integrity(key: dict, book: dict, formulas: dict) -> dict:
    rows = find_sheet(book, key.get('sheet', 'Results'))
    frows = find_sheet(formulas, key.get('sheet', 'Results'))
    cols = result_columns(key, rows)
    if not rows or not cols or len(cols) < 2:
        return {'policy': key.get('formula_policy', 'static_ok'), 'checked': 0, 'verdict': 'not_met' if key.get('formula_policy') == 'live_required' else 'not_applicable', 'why': 'no label/value header to read formulas against'}
    live, literal = ([], [])
    wanted = {norm_label(r['label']) for r in key['rows']}
    for i, r in enumerate(rows[1:], 1):
        if len(r) > max(cols.values()) and r[cols['label']] is not None:
            if result_label(r[cols['label']], wanted) not in wanted:
                continue
            fr = frows[i] if frows and i < len(frows) else None
            (live if fr and len(fr) > cols['value'] and fr[cols['value']] else literal).append(str(r[cols['label']])[:60])
    policy = key.get('formula_policy', 'static_ok')
    out = {'policy': policy, 'checked': len(live) + len(literal), 'live': len(live), 'literal': len(literal), 'literal_labels': literal[:8]}
    if policy == 'live_required':
        out['verdict'] = 'not_met' if literal or not live else 'met'
        out['why'] = 'no graded formula was found' if not live else 'the contract requires a live model; some answers are typed-in values' if literal else 'every graded answer is computed by a formula'
    else:
        out['verdict'] = 'not_applicable'
        out['why'] = 'the contract accepts a static result, so a literal answer is not a defect'
    return out

def score_book(key: dict, book: dict, ai: dict | None=None, formulas: dict | None=None, *, semantic_require_human=True) -> dict:
    verdicts, chosen = score_rows(key, book, formulas)
    verdicts.update(score_tables(key, book))
    verdicts.update(score_checks(key, book))
    for a in key.get('ai_criteria') or []:
        if a['criterion_id'] in verdicts:
            raise ValueError('semantic criterion overlaps a program criterion')
        verdicts[a['criterion_id']] = {'status': 'pending', 'decided_by': 'ai', 'evidence': ''}
    meta = criterion_meta(key)
    if any((m['dimension'] for m in meta.values())):
        for cid, v in verdicts.items():
            v['dimension'] = meta[cid]['dimension']
            v['weight'] = meta[cid]['weight']
    result = {'status': 'SCORED', 'task_id': key.get('task_id'), **aggregate(key, verdicts), 'alternative_sets_applied': chosen, 'verdicts': verdicts}
    if formulas is not None:
        fi = formula_integrity(key, book, formulas)
        result['formula_integrity'] = fi
        if fi.get('verdict') == 'not_met':
            result['critical_status'] = 'FAIL'
            result['accepted'] = 'FAIL'
            result['failed'] = (result.get('failed') or []) + ['hardcoded_answers']
    if result['accepted'] == 'PASS' and any((v.get('status') == 'unresolved' and 'duplicate' in v.get('evidence', '') for v in verdicts.values())):
        result['accepted'] = 'UNRESOLVED'
    if ai is not None:
        from semantic_findings import merge_semantic
        result = merge_semantic(result, key, ai, require_human=semantic_require_human, aggregate_fn=aggregate)
    return result

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--key", required=True)
    parser.add_argument("--candidate", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    key = json.loads(Path(args.key).read_text())
    if key.get("ai_criteria") or key.get("dynamic_checks"):
        raise ValueError("Only declared deterministic criteria are supported")
    output = Path(args.out)
    output.mkdir(parents=True, exist_ok=True)
    answer = Path(args.candidate) / "answer.xlsx"
    if not answer.is_file():
        result = {"status": "OUTPUT_MISSING", "task_id": key.get("task_id"), "score": 0.0}
    else:
        try:
            result = score_book(key, read_static(answer), formulas=read_formulas(answer))
        except Exception as error:
            result = {"status": "JUDGE_ERROR", "task_id": key.get("task_id"), "score": None,
                      "note": f"{type(error).__name__}: {error}"}
    (output / "score.json").write_text(json.dumps(result, ensure_ascii=False, indent=1, default=str) + "\n")
    (output / "reward.txt").write_text(str(result["score"] if result["score"] is not None else "") + "\n")
    print(json.dumps({field: result.get(field) for field in ("status", "score", "accepted")}))
    return 1 if result["status"] == "JUDGE_ERROR" else 0
if __name__ == "__main__":
    sys.exit(main())
