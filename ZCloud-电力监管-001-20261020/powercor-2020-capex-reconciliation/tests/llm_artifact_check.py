"""Validate candidate access evidence without changing scoring criteria."""
import re
from pathlib import Path

import openpyxl


def evidence_errors(details, candidate=Path('/app/output/answer.xlsx')):
    try:
        workbook = openpyxl.load_workbook(candidate, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        sheets = {sheet.title: sheet for sheet in workbook.worksheets}
        target = r'(?:候选(?:交付物|文件|输出文件|工作簿)|(?:/app/output/)?answer\.xlsx|candidate (?:file|workbook|artifact)|output (?:file|workbook|artifact))'
        access = r'(?:不存在|未找到|找不到|缺失|无法(?:读取|打开|解析|访问)|不可(?:读取|读|解析)|missing|unreadable|not found|does not exist|cannot be (?:read|opened|parsed))'
        denied = re.compile(
            target + r'\s*(?:(?:is|was)\s+)?' + access
            + r'|(?:cannot|could not|unable to|failed to) (?:read|parse|open|access) (?:the )?(?:' + target + r'|workbook)'
            + r'|无法(?:解析|读取|打开|访问)\s*' + target
            + r'|(?:文件无法解析|无法解析工作簿)', re.I)
        failures = []
        for entry in (details or {}).values():
            for block in entry if isinstance(entry, list) else [entry]:
                if not isinstance(block, dict) or block.get('kind') != 'agent':
                    continue
                for item in block.get('criteria', []):
                    reason = item.get('reasoning') or ''
                    cited = False
                    for name, sheet in sheets.items():
                        pattern = re.compile(r"(?<![\w])'?" + re.escape(name) + r"'?\s*!\s*\$?([A-Z]{1,3})\$?([1-9][0-9]*)(?![\w])", re.I)
                        for match in pattern.finditer(reason):
                            column = openpyxl.utils.column_index_from_string(match[1].upper())
                            row = int(match[2])
                            if column <= sheet.max_column and row <= sheet.max_row:
                                sheet.cell(row, column).value
                                cited = True
                    problems = []
                    if denied.search(reason):
                        problems.append('candidate_reported_missing_or_unreadable')
                    if not cited:
                        problems.append('candidate_cell_not_cited')
                    if problems:
                        failures.append({'criterion_id': item.get('id', item.get('name')),
                                         'problems': problems, 'reason': reason})
        return failures
    finally:
        workbook.close()
