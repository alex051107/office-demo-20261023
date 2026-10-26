"""Score the required interface, numeric results and independent penalties."""
import copy
import importlib.util
import json
import math
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path
import openpyxl
HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('s3_frozen_bridge', HERE/'frozen_bridge.py')
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)
FROZEN = json.loads((HERE/'frozen_key.json').read_text())

def extract(answer, key):
    evidence = base.extract(answer, FROZEN)
    evidence['interface_errors'] = []
    evidence['header_errors'] = []
    evidence['required_results'] = []
    if evidence['status'] != 'EXTRACTED':
        evidence['interface_errors'].append('required workbook is unavailable')
        return evidence
    workbook = openpyxl.load_workbook(answer, data_only=True, read_only=True)
    try:
        for name, columns in key['required_sheets'].items():
            if name not in workbook.sheetnames:
                evidence['interface_errors'].append(f'required sheet {name!r} is missing or misnamed')
                continue
            sheet = workbook[name]
            header = [cell.value.strip() if isinstance(cell.value, str) else cell.value
                      for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            if header[:len(columns)] != columns:
                evidence['header_errors'].append(f'{name} first columns {header[:len(columns)]!r}; expected {columns!r}')
            for row_index, row in enumerate(sheet.iter_rows(min_row=2), 2):
                if name == key['sheet'] and any(cell.value is not None for cell in row[:4]):
                    cells = list(row[:4])
                    values = [cell.value for cell in cells] + [None] * (4 - len(cells))
                    evidence['required_results'].append({
                        'row': row_index, 'label': values[0], 'value': values[1],
                        'unit': values[2], 'source_note': values[3],
                        'number_format': cells[1].number_format if len(cells) > 1 else 'General',
                    })
    finally:
        workbook.close()
    return evidence


def has_precision(value, number_format, places=3):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        return False
    stored_places = max(0, -Decimal(str(value)).normalize().as_tuple().exponent)
    if value == round(value, places) or stored_places >= places:
        return True
    sections = re.sub(r'"[^"]*"|\\.|_.|\*.', '', number_format or 'General').split(';')
    section = sections[1] if value < 0 and len(sections) > 1 else sections[2] if value == 0 and len(sections) > 2 else sections[0]
    section = re.sub(r'\[[^\]]*\]', '', section).strip()
    if section.casefold() == 'general':
        return stored_places >= places
    decimals = re.search(r'[0#?]+\.([0#?]+)', section)
    if not decimals:
        return False
    digits = decimals.group(1)
    return max(digits.count('0'), min(stored_places, len(digits))) >= places


def contract_penalties(evidence, key):
    rows = evidence['required_results']
    expected = [item['label'] for item in key['rows']]
    first = {label: next((row for row in rows if row['label'] == label), {}) for label in expected}
    source_bad = [label for label, row in first.items() if not isinstance(row.get('source_note'), str)
                  or not re.search(r'[^\W\d_]', row['source_note'], re.UNICODE)]
    unit_bad = [item['label'] for item in key['rows'] if first[item['label']].get('unit') != item['unit']]
    sequence = [row['label'] for row in rows if row['label'] in expected]
    counts = Counter(row['label'] for row in rows if row['label'] in expected)
    extra = [row['row'] for row in rows if row['label'] not in expected
             and row['value'] is not None and str(row['value']).strip()]
    duplicates = [label for label, count in counts.items() if count > 1]
    precision_bad = [label for label, row in first.items()
                     if not has_precision(row.get('value'), row.get('number_format'))]
    return {
        'S01': (bool(evidence['header_errors']), '; '.join(evidence['header_errors'])
                or 'required leading header columns match after trimming spaces'),
        'S02': (bool(unit_bad), f'units not exactly as specified: {unit_bad!r}'),
        'S03': (bool(source_bad), f'missing or non-textual source references: {source_bad!r}'),
        'S04': (sequence != expected or [first[label].get('row') for label in expected] != [2, 3, 4],
                f'required row sequence: {sequence!r}; expected rows 2, 3, 4'),
        'S05': (bool(extra or duplicates), f'unrequested business rows: {extra!r}; duplicate labels: {duplicates!r}'),
        'S06': (bool(precision_bad), f'results not showing at least three decimal places: {precision_bad!r}'),
    }


def score_policy(evidence, key, verdicts):
    reasons = list(evidence.get('header_errors', []))
    for criterion_id in ['S02', 'S03']:
        if verdicts[criterion_id]['credit']:
            reasons.append(verdicts[criterion_id]['evidence'])
    expected = {item['label'] for item in key['rows']}
    fabricated = [row['row'] for row in evidence['required_results'] if row['label'] not in expected
                  and row['value'] is not None and str(row['value']).strip()]
    if fabricated:
        reasons.append(f'fabricated extra business records in rows {fabricated!r}')
    return {
        'required_sheet_gate': {'triggered': bool(evidence['interface_errors']),
                                'reasons': evidence['interface_errors'], 'maximum_score': 0.0},
        'must_clause_cap': {'triggered': bool(reasons), 'reasons': reasons, 'maximum_score': 0.69},
    }


def score(evidence, key):
    all_c = key.get('rows', []) + key.get('tables', []) + key.get('checks', [])
    if evidence['status'] == 'JUDGE_ERROR':
        return base.score(evidence, FROZEN)
    errors = evidence.get('interface_errors', [])
    if errors:
        verdicts = {c['criterion_id']: {
            'credit': int(c['criterion_id'] == 'S01'), 'decided_by': 'program',
            'status': 'triggered' if c['criterion_id'] == 'S01' else 'gated',
            'evidence': '; '.join(errors),
        } for c in all_c}
    else:
        numeric_evidence = copy.deepcopy(evidence)
        numeric_rows = numeric_evidence['book'][key['sheet']]
        if numeric_rows:
            columns = key['required_sheets'][key['sheet']]
            numeric_rows[0] = list(columns) + numeric_rows[0][len(columns):]
        original = base.score(numeric_evidence, FROZEN)
        verdicts = {c['criterion_id']: copy.deepcopy(original['verdicts'][c['criterion_id']])
                    for c in all_c if c['criterion_id'] in original['verdicts']}
        for criterion_id, (violated, reason) in contract_penalties(evidence, key).items():
            verdicts[criterion_id] = {'credit': int(violated), 'decided_by': 'program',
                                     'status': 'triggered' if violated else 'not_triggered', 'evidence': reason}
    total = sum(c['weight'] for c in all_c if c['weight']>0)
    numerator = sum(c['weight']*(1-round(1-verdicts[c['criterion_id']]['credit'],4) if c['weight']<0 else round(verdicts[c['criterion_id']]['credit'],4)) for c in all_c)
    pooled = max(0, min(1, numerator / total))
    policy = score_policy(evidence, key, verdicts)
    final_score = 0.0 if errors else min(pooled, 0.69) if policy['must_clause_cap']['triggered'] else pooled
    return {'status': evidence['status'], 'score': round(final_score, 6),
            'signed_pooled_score': round(pooled, 6), 'score_policy': policy,
            'interface_gate': {'passed': not errors, 'reasons': errors}, 'verdicts':verdicts}
