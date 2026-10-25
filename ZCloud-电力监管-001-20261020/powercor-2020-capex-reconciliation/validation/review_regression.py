"""Offline regression workbooks for the Powercor interface and scoring contract."""
import importlib.util
import json
import math
import sys
from pathlib import Path

import openpyxl

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def evaluate(candidate):
    scorer = load_module('regression_scorer', ROOT / 'tests/score_output.py')
    key = json.loads((ROOT / 'tests/results_key.json').read_text())
    evidence = scorer.extract(candidate, key)
    return scorer.score(evidence, key)


def mutate(workbook, name):
    sheet = workbook['Results']
    if name == 'sheet_renamed':
        sheet.title = 'WrongResults'
    elif name == 'sheet_case':
        sheet.title = 'results'
    elif name == 'sheet_missing':
        workbook.remove(sheet)
    elif name == 'header_reversed':
        for row in sheet:
            values = [cell.value for cell in row][::-1]
            for cell, value in zip(row, values):
                cell.value = value
    elif name == 'extra_column':
        sheet['E1'] = 'extra'
    elif name == 'extra_column_no_header':
        sheet['E2'] = 'extra'
    elif name == 'source_header_missing':
        sheet['D1'] = None
    elif name == 'header_case':
        sheet['A1'] = 'Label'
    elif name == 'header_shifted':
        sheet.insert_rows(1)
    elif name == 'rows_reversed':
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2, max_row=4)][::-1]
        for row_index, values in enumerate(rows, 2):
            for column, value in enumerate(values, 1):
                sheet.cell(row_index, column, value)
    elif name == 'rows_shifted':
        sheet.insert_rows(2)
    elif name in {'sources_empty', 'source_one_empty', 'sources_numeric'}:
        for row_index in range(2, 5 if name != 'source_one_empty' else 3):
            sheet.cell(row_index, 4).value = 123 if name == 'sources_numeric' else None
    elif name in {'unit_wrong', 'unit_case'}:
        for row_index in range(2, 5):
            sheet.cell(row_index, 3, 'USD thousand real' if name == 'unit_wrong' else 'aud million nominal')
    elif name == 'extra_record':
        sheet.append(['Unrequested asset reconciliation', 999999, 'AUD million nominal', 'RFM input'])
    elif name == 'extra_text_record':
        sheet.append(['Unrequested asset reconciliation', 'not numeric', 'AUD million nominal', 'RFM input'])
    elif name == 'duplicate_record':
        sheet.append([cell.value for cell in sheet[2]])
    elif name == 'orphan_value':
        sheet.append([None, 999999, 'AUD million nominal', 'RFM input'])
    elif name in {'rounded_two_general', 'rounded_two_formatted', 'rounded_three'}:
        digits = 3 if name == 'rounded_three' else 2
        for row_index in range(2, 5):
            cell = sheet.cell(row_index, 2)
            cell.value = round(cell.value, digits)
            cell.number_format = 'General' if name == 'rounded_two_general' else '0.000'
    elif name == 'two_decimal_display':
        for row_index in range(2, 5):
            sheet.cell(row_index, 2).number_format = '0.00'
    elif name in {'inside_tolerance', 'outside_tolerance'}:
        sheet['B2'] = sheet['B2'].value + (0.00049 if name == 'inside_tolerance' else 0.00051)
    elif name == 'full_precision_general':
        for row_index in range(2, 5):
            sheet.cell(row_index, 2).number_format = 'General'
    elif name == 'extra_support_sheet':
        workbook.create_sheet('Unconstrained support').append(['arbitrary', 'support layout', 12])
    elif name != 'golden':
        raise ValueError(name)


def independent_calculation():
    inputs = openpyxl.load_workbook(ROOT / 'environment/input_files/Powercor_RFM_inputs.xlsx', data_only=True)
    sheet = inputs['RFM input']
    rate = (1 + sheet['F287'].value) / (1 + sheet['F284'].value) * (1 + sheet['F282'].value) - 1
    factor = math.sqrt(1 + rate)
    totals = {'network': 0.0, 'non-network': 0.0, 'all': 0.0}
    groups = {7: 'network', 8: 'network', 10: 'network', 11: 'non-network', 12: 'non-network', 13: 'non-network'}
    golden = openpyxl.load_workbook(ROOT / 'solution/golden_output/answer.xlsx', data_only=True)
    errors = []
    count = 0
    for source_row in range(7, 57):
        if not isinstance(sheet.cell(source_row, 6).value, str):
            continue
        actual = sum(coef * (sheet.cell(source_row + offset, 6).value or 0)
                     for offset, coef in [(54, 1), (108, -1), (162, -1)])
        forecast = sheet.cell(source_row, 12).value or 0
        difference = actual * factor - forecast
        count += 1
        errors.append(abs(golden['Working'].cell(count + 1, 5).value - difference))
        totals['all'] += difference
        if source_row in groups:
            totals[groups[source_row]] += difference
    for row_index, group in enumerate(['network', 'non-network', 'all'], 2):
        errors.append(abs(golden['Results'].cell(row_index, 2).value - totals[group]))
    inputs.close()
    golden.close()
    assert count == 12 and max(errors) < 1e-12
    return {'asset_classes': count, 'applicable_rate': rate, 'year_end_factor': factor,
            'totals': totals, 'max_absolute_difference': max(errors)}


def main():
    folder = ROOT / 'validation/negative_cases'
    folder.mkdir(parents=True, exist_ok=True)
    cases = [
        ('sheet_renamed', 'Results改错表名', 0.0, ['S01'], True),
        ('header_reversed', 'Results四列逆序', 0.9, ['S01'], True),
        ('unit_wrong', '单位改为USD thousand real', 0.9, ['S02'], False),
        ('extra_column', 'Results增加第五列表头', 1.0, ['S01'], True),
        ('sources_empty', '清空三条source_note', 1.0, ['S03'], False),
        ('rows_reversed', 'Results三条结果逆序', 1.0, ['S04'], False),
        ('rounded_two_general', '三项四舍五入两位且为General', 1.0, ['S06'], False),
        ('sheet_case', 'Results仅表名大小写错误', None, ['S01'], True),
        ('sheet_missing', '删除Results表', None, ['S01'], True),
        ('extra_column_no_header', '第五列有数据但表头为空', None, ['S01'], True),
        ('source_header_missing', 'source_note表头清空', None, ['S01'], True),
        ('header_case', 'label表头大小写错误', None, ['S01'], True),
        ('header_shifted', '表头移至第二行', None, ['S01'], True),
        ('rows_shifted', '必需结果移至第3至5行', None, ['S04'], False),
        ('source_one_empty', '仅一条source_note为空', None, ['S03'], False),
        ('sources_numeric', '来源说明全部换成纯数字', None, ['S03'], False),
        ('unit_case', '单位大小写与题面不一致', None, ['S02'], False),
        ('extra_record', '尾部新增未请求的数字业务行', None, ['S05'], False),
        ('extra_text_record', '尾部新增未请求的文字业务行', None, ['S05'], False),
        ('duplicate_record', '重复首条必需业务行', None, ['S04', 'S05'], False),
        ('orphan_value', '无标签但有金额的额外业务行', None, ['S05'], False),
        ('rounded_two_formatted', '两位存储值以0.000显示', None, [], False),
        ('two_decimal_display', '保留完整数值但只显示两位', None, ['S06'], False),
        ('outside_tolerance', '第一项增加0.00051超出容差', None, [], False),
    ]
    rows = []
    for name, mutation, old_score, penalties, gate in cases:
        workbook = openpyxl.load_workbook(ROOT / 'solution/golden_output/answer.xlsx')
        mutate(workbook, name)
        path = folder / f'{name}.xlsx'
        workbook.save(path)
        workbook.close()
        result = evaluate(path)
        triggered = [criterion_id for criterion_id, verdict in result['verdicts'].items()
                     if criterion_id[0] in {'S', 'P'} and verdict['credit']]
        assert result['score'] < 1.0, name
        assert not gate or result['score'] == 0.0, name
        assert set(penalties) <= set(triggered), (name, triggered)
        rows.append({'case': name, 'workbook': path.name, 'mutation': mutation,
                     'expected_outcome': 'interface_gate_0' if gate else 'penalty_or_numeric_loss',
                     'expected_penalties': penalties, 'old_review_program_score': old_score,
                     'old_score_source': 'review section 三·A' if old_score is not None else '未在原审查中运行',
                     'new_program_score': result['score'], 'penalties_triggered': triggered,
                     'numeric_criteria_failed': [criterion_id for criterion_id, verdict in result['verdicts'].items()
                                                 if criterion_id.startswith('R') and not verdict['credit']]})
    positives = []
    controls = ROOT / 'validation/positive_cases'
    controls.mkdir(parents=True, exist_ok=True)
    for name in ['golden', 'full_precision_general', 'rounded_three', 'inside_tolerance', 'extra_support_sheet']:
        workbook = openpyxl.load_workbook(ROOT / 'solution/golden_output/answer.xlsx')
        mutate(workbook, name)
        path = controls / f'{name}.xlsx'
        workbook.save(path)
        workbook.close()
        result = evaluate(path)
        assert result['score'] == 1.0, (name, result)
        positives.append({'case': name, 'program_score': result['score']})
    report = {'scorer_version': 'v2 (2026-10-26)', 'scope': 'programmatic subset only; no new LLM judgments',
              'negative_cases': rows, 'positive_controls': positives, 'independent_calculation': independent_calculation()}
    (folder / 'cases.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'negative_cases': len(rows), 'interface_cases': sum(row['expected_outcome'] == 'interface_gate_0' for row in rows),
                      'positive_controls': len(positives), 'independent_calculation': report['independent_calculation']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
