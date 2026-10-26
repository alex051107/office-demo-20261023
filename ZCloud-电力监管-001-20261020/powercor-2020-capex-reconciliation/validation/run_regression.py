"""Read-only offline regression entry for the task and its saved report."""
import argparse
import hashlib
import importlib.util
import json
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True

"""Offline regression workbooks for the Powercor interface and scoring contract."""
import importlib.util
import json
import math
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill
sys.dont_write_bytecode = True
scoring_ROOT = Path(__file__).resolve().parents[1]

def scoring_load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def scoring_evaluate(candidate):
    scorer = scoring_load_module('regression_scorer', scoring_ROOT / 'tests/score_output.py')
    key = json.loads((scoring_ROOT / 'tests/results_key.json').read_text())
    evidence = scorer.extract(candidate, key)
    return scorer.score(evidence, key)

def scoring_mutate(workbook, name):
    sheet = workbook['Results']
    if name == 'sheet_renamed':
        sheet.title = 'WrongResults'
    elif name == 'sheet_case':
        sheet.title = 'results'
    elif name == 'sheet_missing':
        workbook.remove(sheet)
    elif name == 'header_reversed':
        for row in sheet:
            values = [cell.value for cell in row][::-1]
            for cell, value in zip(row, values):
                cell.value = value
    elif name == 'extra_column':
        sheet['E1'] = 'extra'
    elif name == 'extra_column_no_header':
        sheet['E2'] = 'extra'
    elif name == 'source_header_missing':
        sheet['D1'] = None
    elif name == 'header_case':
        sheet['A1'] = 'Label'
    elif name == 'header_order_only':
        sheet['A1'], sheet['B1'] = ('value', 'label')
    elif name == 'header_spaces':
        for cell in sheet[1]:
            cell.value = f'  {cell.value}  '
    elif name == 'empty_styled_column':
        sheet['E1'].fill = PatternFill('solid', fgColor='FFEECC')
    elif name == 'label_only_note':
        sheet['A6'] = 'Supporting note without a business value'
    elif name == 'header_shifted':
        sheet.insert_rows(1)
    elif name == 'rows_reversed':
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2, max_row=4)][::-1]
        for row_index, values in enumerate(rows, 2):
            for column, value in enumerate(values, 1):
                sheet.cell(row_index, column, value)
    elif name == 'rows_shifted':
        sheet.insert_rows(2)
    elif name in {'sources_empty', 'source_one_empty', 'sources_numeric', 'sources_numeric_text'}:
        for row_index in range(2, 5 if name != 'source_one_empty' else 3):
            sheet.cell(row_index, 4).value = 123 if name == 'sources_numeric' else '123.45' if name == 'sources_numeric_text' else None
    elif name in {'unit_wrong', 'unit_case'}:
        for row_index in range(2, 5):
            sheet.cell(row_index, 3, 'USD thousand real' if name == 'unit_wrong' else 'aud million nominal')
    elif name == 'extra_record':
        sheet.append(['Unrequested asset reconciliation', 999999, 'AUD million nominal', 'RFM input'])
    elif name == 'extra_text_record':
        sheet.append(['Unrequested asset reconciliation', 'not numeric', 'AUD million nominal', 'RFM input'])
    elif name == 'duplicate_record':
        sheet.append([cell.value for cell in sheet[2]])
    elif name == 'orphan_value':
        sheet.append([None, 999999, 'AUD million nominal', 'RFM input'])
    elif name in {'rounded_two_general', 'rounded_two_formatted', 'rounded_three'}:
        digits = 3 if name == 'rounded_three' else 2
        for row_index in range(2, 5):
            cell = sheet.cell(row_index, 2)
            cell.value = round(cell.value, digits)
            cell.number_format = 'General' if name == 'rounded_two_general' else '0.000'
    elif name == 'two_decimal_display':
        for row_index in range(2, 5):
            sheet.cell(row_index, 2).number_format = '0.00'
    elif name == 'nonnumeric_value':
        sheet['B2'] = '13.690'
    elif name == 'rounded_three_general':
        for row_index in range(2, 5):
            sheet.cell(row_index, 2).value = round(sheet.cell(row_index, 2).value, 3)
            sheet.cell(row_index, 2).number_format = 'General'
    elif name in {'inside_tolerance', 'outside_tolerance'}:
        sheet['B2'] = sheet['B2'].value + (0.00149 if name == 'inside_tolerance' else 0.00151)
    elif name == 'rounded_asset_aggregation':
        assets = list(workbook['Working'].iter_rows(min_row=2, max_row=13))
        for row_index, group in [(2, 'network'), (3, 'non-network'), (4, None)]:
            sheet.cell(row_index, 2).value = sum((round(row[4].value, 3) for row in assets if group is None or row[1].value == group))
    elif name == 'full_precision_general':
        for row_index in range(2, 5):
            sheet.cell(row_index, 2).number_format = 'General'
    elif name == 'extra_support_sheet':
        workbook.create_sheet('Unconstrained support').append(['arbitrary', 'support layout', 12])
    elif name != 'golden':
        raise ValueError(name)

def scoring_independent_calculation():
    inputs = openpyxl.load_workbook(scoring_ROOT / 'environment/input_files/Powercor_RFM_inputs.xlsx', data_only=True)
    sheet = inputs['RFM input']
    rate = (1 + sheet['F287'].value) / (1 + sheet['F284'].value) * (1 + sheet['F282'].value) - 1
    factor = math.sqrt(1 + rate)
    totals = {'network': 0.0, 'non-network': 0.0, 'all': 0.0}
    groups = {7: 'network', 8: 'network', 10: 'network', 11: 'non-network', 12: 'non-network', 13: 'non-network'}
    golden = openpyxl.load_workbook(scoring_ROOT / 'solution/golden_output/answer.xlsx', data_only=True)
    errors = []
    count = 0
    for source_row in range(7, 57):
        if not isinstance(sheet.cell(source_row, 6).value, str):
            continue
        actual = sum((coef * (sheet.cell(source_row + offset, 6).value or 0) for offset, coef in [(54, 1), (108, -1), (162, -1)]))
        forecast = sheet.cell(source_row, 12).value or 0
        difference = actual * factor - forecast
        count += 1
        errors.append(abs(golden['Working'].cell(count + 1, 5).value - difference))
        totals['all'] += difference
        if source_row in groups:
            totals[groups[source_row]] += difference
    for row_index, group in enumerate(['network', 'non-network', 'all'], 2):
        errors.append(abs(golden['Results'].cell(row_index, 2).value - totals[group]))
    inputs.close()
    golden.close()
    assert count == 12 and max(errors) < 1e-12
    return {'asset_classes': count, 'applicable_rate': rate, 'year_end_factor': factor, 'totals': totals, 'max_absolute_difference': max(errors)}

def scoring_main(temporary):
    folder = Path(temporary) / 'negative_cases'
    folder.mkdir(parents=True, exist_ok=True)
    cases = [('sheet_renamed', 'Results改错表名', 0.0, ['S01'], True), ('header_reversed', 'Results四列逆序', 0.9, ['S01'], False), ('unit_wrong', '单位改为USD thousand real', 0.9, ['S02'], False), ('sources_empty', '清空三条source_note', 1.0, ['S03'], False), ('rows_reversed', 'Results三条结果逆序', 1.0, ['S04'], False), ('rounded_two_general', '三项四舍五入两位且为General；存储精度合规但数值误差失分', 1.0, [], False), ('sheet_case', 'Results仅表名大小写错误', None, ['S01'], True), ('sheet_missing', '删除Results表', None, ['S01'], True), ('source_header_missing', 'source_note表头清空', None, ['S01'], False), ('header_case', 'label表头大小写错误', None, ['S01'], False), ('header_order_only', '仅交换前两个表头，数值内容不变', None, ['S01'], False), ('header_shifted', '表头移至第二行', None, ['S01'], False), ('rows_shifted', '必需结果移至第3至5行', None, ['S04'], False), ('source_one_empty', '仅一条source_note为空', None, ['S03'], False), ('sources_numeric', '来源说明全部换成纯数字', None, ['S03'], False), ('sources_numeric_text', '来源说明全部换成纯数字字符串', None, ['S03'], False), ('unit_case', '单位大小写与题面不一致', None, ['S02'], False), ('extra_record', '尾部新增未请求的数字业务行', None, ['S05'], False), ('extra_text_record', '尾部新增未请求的文字业务行', None, ['S05'], False), ('duplicate_record', '重复首条必需业务行', None, ['S04', 'S05'], False), ('orphan_value', '无标签但有金额的额外业务行', None, ['S05'], False), ('rounded_two_formatted', '两位存储值以0.000显示', None, [], False), ('nonnumeric_value', '一条结果换成文本而非数值', None, ['S06'], False), ('outside_tolerance', '第一项增加0.00151超出该派生行0.0015容差', None, [], False)]
    rows = []
    cap_cases = {'header_reversed', 'source_header_missing', 'header_case', 'header_order_only', 'header_shifted', 'sources_empty', 'source_one_empty', 'sources_numeric', 'sources_numeric_text', 'unit_wrong', 'unit_case', 'extra_record', 'extra_text_record', 'orphan_value'}
    for name, mutation, old_score, penalties, gate in cases:
        workbook = openpyxl.load_workbook(scoring_ROOT / 'solution/golden_output/answer.xlsx')
        scoring_mutate(workbook, name)
        path = folder / f'{name}.xlsx'
        workbook.save(path)
        workbook.close()
        result = scoring_evaluate(path)
        triggered = [criterion_id for criterion_id, verdict in result['verdicts'].items() if criterion_id[0] in {'S', 'P'} and verdict['credit']]
        assert result['score'] < 1.0, name
        assert not gate or result['score'] == 0.0, name
        if name in cap_cases:
            assert result['score_policy']['must_clause_cap']['triggered'], name
            assert result['score'] <= 0.69, name
            if name in {'header_order_only', 'header_case', 'source_header_missing'}:
                assert result['score'] == 0.69, name
                assert all((result['verdicts'][row['criterion_id']]['credit'] for row in json.loads((scoring_ROOT / 'tests/results_key.json').read_text())['rows'])), name
        elif not gate:
            assert not result['score_policy']['must_clause_cap']['triggered'], name
        assert set(penalties) <= set(triggered), (name, triggered)
        rows.append({'case': name, 'workbook': path.name, 'mutation': mutation, 'expected_outcome': 'interface_gate_0' if gate else 'must_clause_cap_0.69' if name in cap_cases else 'penalty_below_1', 'expected_penalties': penalties, 'old_review_program_score': old_score, 'old_score_source': 'review section 三·A' if old_score is not None else '未在原审查中运行', 'new_program_score': result['score'], 'penalties_triggered': triggered, 'signed_pooled_score': result['signed_pooled_score'], 'score_policy': result['score_policy'], 'numeric_criteria_failed': [criterion_id for criterion_id, verdict in result['verdicts'].items() if criterion_id.startswith('R') and (not verdict['credit'])]})
    positives = []
    controls = Path(temporary) / 'positive_cases'
    controls.mkdir(parents=True, exist_ok=True)
    for name in ['golden', 'full_precision_general', 'rounded_three', 'rounded_three_general', 'rounded_asset_aggregation', 'inside_tolerance', 'extra_support_sheet', 'header_spaces', 'empty_styled_column', 'extra_column', 'extra_column_no_header', 'two_decimal_display', 'label_only_note']:
        workbook = openpyxl.load_workbook(scoring_ROOT / 'solution/golden_output/answer.xlsx')
        scoring_mutate(workbook, name)
        path = controls / f'{name}.xlsx'
        workbook.save(path)
        workbook.close()
        result = scoring_evaluate(path)
        assert result['score'] == 1.0, (name, result)
        positives.append({'case': name, 'program_score': result['score']})
    for name in ['extra_column', 'extra_column_no_header', 'two_decimal_display']:
        (folder / f'{name}.xlsx').unlink(missing_ok=True)
    report = {'scorer_version': 'v2 (2026-10-26)', 'scope': 'programmatic subset only; no new LLM judgments', 'negative_cases': rows, 'positive_controls': positives, 'independent_calculation': scoring_independent_calculation()}
    report['precision_rule'] = 'p decimals: a matching number_format or stored value equal to round(value,p); at least p also accepts stored decimal places >= p. Numeric tolerance is separate.'
    (folder / 'cases.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'negative_cases': len(rows), 'interface_cases': sum((row['expected_outcome'] == 'interface_gate_0' for row in rows)), 'positive_controls': len(positives), 'independent_calculation': report['independent_calculation']}, ensure_ascii=False))

"""Check post-pooling policies through the delivered finalizer without model calls."""
import importlib.util
import json
import subprocess
import sys
import tempfile
from pathlib import Path
sys.dont_write_bytecode = True
pooling_ROOT = Path(__file__).resolve().parents[1]

def pooling_load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def pooling_main():
    scorer = pooling_load_module('policy_scorer', pooling_ROOT / 'tests/score_output.py')
    finalizer = pooling_load_module('policy_finalizer', pooling_ROOT / 'tests/finalize.py')
    key = json.loads((pooling_ROOT / 'tests/results_key.json').read_text())
    criteria = key['rows'] + key['checks']
    outcomes = []
    for candidate in sorted((pooling_ROOT / 'validation').glob('*_cases/*.xlsx')):
        result = scorer.score(scorer.extract(candidate, key), key)
        entries = [{'name': item['criterion_id'], 'weight': abs(item['weight']), 'negate': item['weight'] < 0, 'value': 1 - result['verdicts'][item['criterion_id']]['credit'] if item['weight'] < 0 else result['verdicts'][item['criterion_id']]['credit']} for item in criteria]
        entries += [{'name': name, 'weight': 3, 'negate': True, 'value': 1} for name in ['L01', 'L02', 'L03']]
        details = {'reward': {'kind': 'programmatic_policy_control', 'criteria': entries}, 'score_policy': result['score_policy']}
        pooled, counted, broken = finalizer.pooled_score(details)
        assert counted == 13 and broken == 0
        assert round(finalizer.apply_score_policy(pooled, result['score_policy']), 6) == result['score']
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            (folder / 'reward.json').write_text(json.dumps({'reward': pooled}))
            (folder / 'reward-details.json').write_text(json.dumps(details))
            subprocess.run([sys.executable, str(pooling_ROOT / 'tests/finalize.py'), '--graded', str(folder / 'reward.json'), '--graded-rc', '0', '--out', str(folder / 'final/reward.json')], check=True)
            receipt = json.loads((folder / 'final/reward.json').read_text())
            assert receipt['reward'] == result['score'] and (not receipt['verifier_error']), candidate.name
            recorded = json.loads((folder / 'final/reward-details.json').read_text())
            assert recorded['score_policy'] == result['score_policy']
        outcomes.append({'case': candidate.stem, 'score': result['score'], 'finalizer_cli_passed': True})
    low = {'required_sheet_gate': {'triggered': False}, 'must_clause_cap': {'triggered': True}}
    assert finalizer.apply_score_policy(0.4, low) == 0.4
    low['required_sheet_gate']['triggered'] = True
    assert finalizer.apply_score_policy(1.0, low) == 0.0
    for value in [1, 1.2, 1.23, 1.234, 1.234567]:
        assert scorer.has_precision(value, 'General')
    assert scorer.has_precision(1.234567, '0.00')
    assert not scorer.has_precision('1.234', '0.000')
    report = {'status': 'PASS', 'finalizer_cli_cases': len(outcomes), 'cap_never_increases_score': True, 'gate_overrides_full_pool': True, 'precision_controls_passed': True, 'cases': outcomes}
    print(json.dumps({key: value for key, value in report.items() if key != 'cases'}))

"""Regression checks for candidate-access evidence and saved judgments."""
import argparse
import copy
import importlib.util
import json
import subprocess
import sys
import tempfile
from pathlib import Path
access_TASK = Path(__file__).resolve().parents[1]

def access_load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def access_main():
    finalizer = access_load('artifact_finalizer', access_TASK / 'tests/finalize.py')
    checker = access_load('artifact_checker', access_TASK / 'tests/llm_artifact_check.py')
    golden = access_TASK / 'solution/golden_output/answer.xlsx'
    blind = json.loads((access_TASK / 'history/validation/judge_fixtures/saved_blind.json').read_text())
    with tempfile.TemporaryDirectory() as temporary:
        folder = Path(temporary)
        graded = folder / 'graded/reward.json'
        graded.parent.mkdir()
        graded.write_text('{"programmatic": 1.0}')
        valid = {'reward': [{'kind': 'programmatic', 'criteria': [{'name': 'control', 'weight': 10, 'value': 1.0, 'negate': False}]}, {'kind': 'agent', 'criteria': [{'id': name, 'weight': 3, 'value': 1.0, 'negate': True, 'reasoning': 'Results!D2 contains a source citation with no unsupported assertion.'} for name in ['L01', 'L02', 'L03']]}]}
        cases = [('evidence_backed_control', valid, False), ('saved_blind_C_golden_r4', {'reward': [valid['reward'][0], blind['agent']]}, True)]
        for name, reason, expected in [('no_cell', 'No unsupported assertions were found.', True), ('source_only_cell', 'The source Text!A131 supports this conclusion.', True), ('nonexistent_candidate_cell', 'Results!ZZ9999 supports this conclusion.', True), ('missing_with_cell', 'Results!D2 cannot be read because candidate workbook is missing.', True), ('unreadable_with_cell', 'Results!D2：无法解析候选文件。', True), ('missing_source_not_artifact', '候选 Results!D2 引用了不存在的来源文件。', False), ('missing_source_english', 'The candidate workbook Results!D2 cites a missing source file.', False), ('quoted_cell', "'Results'!$D$2 contains no audit-completion claim.", False)]:
            detail = copy.deepcopy(valid)
            detail['reward'][1]['criteria'][0]['reasoning'] = reason
            cases.append((name, detail, expected))
        for name, detail, expected_error in cases:
            graded.with_name('reward-details.json').write_text(json.dumps(detail))
            output = folder / name / 'reward.json'
            subprocess.run([sys.executable, str(access_TASK / 'tests/finalize.py'), '--graded', str(graded), '--graded-rc', '0', '--candidate', str(golden), '--out', str(output)], check=True)
            receipt = json.loads(output.read_text())
            assert bool(receipt['verifier_error']) == expected_error, (name, receipt)
            if expected_error:
                assert receipt['reward'] == 0
                exit_message = json.loads(output.with_name('reward_exit_message.json').read_text())
                assert exit_message['exit_code'] == 'judge:artifact_access_error', name
            else:
                assert receipt['reward'] == 1.0, name
                assert not output.with_name('reward_exit_message.json').exists(), name
        saved_valid = 0
        for fixture in (access_TASK / 'history/validation/judge_fixtures').glob('saved_valid_*.json'):
            payload = json.loads(fixture.read_text())
            assert checker.evidence_errors({'reward': payload['agent']}, golden) == [], fixture.name
            saved_valid += 1
        assert checker.evidence_errors({'reward': blind['agent']}, folder / 'missing.xlsx') == []
    print(json.dumps({'artifact_access_regressions': len(cases), 'saved_blind_cases': 1, 'saved_valid_cases': saved_valid, 'missing_candidate_keeps_program_gate': True}))


def check_saved_cases(task):
    engine = import_module('case_inventory_engine', task / 'tests/score_output.py')
    key = json.loads((task / 'tests/results_key.json').read_text())
    counts = {}
    for category in ('negative_cases', 'positive_cases'):
        folder = task / 'validation' / category
        cases = json.loads((folder / 'cases.json').read_text())['cases']
        assert {entry['file'] for entry in cases} == {path.name for path in folder.glob('*.xlsx')}
        for entry in cases:
            result = engine.score(engine.extract(folder / entry['file'], key), key)
            expected = entry.get('new_program_score', entry.get('program_score', entry.get('expected_program_score')))
            assert expected is not None and abs(result['score'] - expected) < 0.000001, entry['file']
            assert result['score'] == 1.0 if category == 'positive_cases' else result['score'] < 1.0
        counts[category] = len(cases)
    return counts


def check_saved_report(task, track):
    import statistics

    def read(path):
        return json.loads(path.read_text())

    engine = import_module('saved_report_engine', task / 'tests/score_output.py')
    finalizer = import_module('saved_report_finalizer', task / 'tests/finalize.py')
    key = read(task / 'tests/results_key.json')
    golden = track / 'golden/candidate.xlsx'
    assert golden.read_bytes() == (task / 'solution/golden_output/answer.xlsx').read_bytes()
    assert engine.score(engine.extract(golden, key), key)['score'] == 1.0
    summary = read(track / 'rejudge_summary.json')
    assert summary == read(task / 'validation/rejudge_summary.json')
    consistency = read(track / 'consistency.json')
    group = consistency['groups'][0]
    judges = []
    expected_ids = {item['id'] for item in read(task / 'rubrics.json')['items']}
    for identifier in group['run_ids']:
        folder = track / 'judge_runs' / identifier
        reward = read(folder / 'reward.json')
        details = read(folder / 'reward-details.json')
        judge = read(folder / 'judge.json')
        assert reward['verifier_error'] == judge['verifier_error'] == 0
        assert finalizer.llm_evidence_errors(details, golden) == []
        pooled, counted, broken = finalizer.pooled_score(details)
        if hasattr(finalizer, 'apply_post_pool_rules'):
            score = finalizer.apply_post_pool_rules(details, pooled)
        elif hasattr(finalizer, 'apply_score_policy'):
            score = finalizer.apply_score_policy(pooled, details['score_policy'])
        else:
            score = finalizer.apply_score_rules(pooled, finalizer.program_result(golden))
        assert not broken and counted == len(expected_ids)
        assert reward['reward'] == judge['score_final'] == round(score, 6)
        assert {item['criterion_id'] for item in judge['criteria']} == expected_ids
        assert judge['reward_details_sha256'] == hashlib.sha256((folder / 'reward-details.json').read_bytes()).hexdigest()
        assert judge['judge_model'] == summary['judge_model'] == 'claude-sonnet-5-5'
        assert [block['judge']['model'] for block in details['reward'] if block.get('kind') == 'agent'] == [judge['judge_model']]
        assert judge['source_run'] == identifier
        judges.append(judge)
    scores = [judge['score_final'] for judge in judges]
    sigma = statistics.stdev(scores) if len(scores) > 1 else None
    assert group['n'] == summary['valid_runs'] == len(scores)
    assert group['total_score']['values'] == scores
    assert group['total_score']['stddev_sample'] == summary['total_sigma_sample'] == sigma
    assert len(group['criteria']) == len(expected_ids)
    for criterion in group['criteria']:
        values = [next(item['score'] for item in judge['criteria'] if item['criterion_id'] == criterion['criterion_id']) for judge in judges]
        deviation = statistics.stdev(values) if len(values) > 1 else None
        assert criterion['values'] == values and criterion['n'] == len(values)
        assert criterion['stddev_sample'] == summary['criterion_sigma_sample'][criterion['criterion_id']] == deviation
        assert criterion['pass_stddev_lt_0_05'] == (deviation is not None and deviation < 0.05)
    requirement = len(scores) == 5 and sigma is not None and sigma < 0.05 and all(item['pass_stddev_lt_0_05'] for item in group['criteria'])
    assert group['requirement_met'] == summary['consistency_requirement_met'] == requirement
    reward = read(track / 'golden/reward.json')
    judge = read(track / 'golden/judge.json')
    if scores:
        assert judge['source_run'] == summary['representative'] == judges[0]['source_run']
        assert reward['reward'] == judge['score_final'] == scores[0] == summary['new_score']
        assert reward['score_available'] and judge['judge_completed']
        assert (track / 'golden/reward-details.json').read_bytes() == (track / 'judge_runs' / group['run_ids'][0] / 'reward-details.json').read_bytes()
    else:
        assert reward['verifier_error'] == 1 and reward['score_available'] is False
        assert judge['score_final'] is None and judge['judge_completed'] is False
        assert not list((track / 'judge_runs').glob('golden_r[1-5]'))
    quality = read(track / 'quality_report.json')
    partial = read(track / 'discrimination_cases.json')
    assert quality['target']['score_final'] == quality['discrimination']['reference_score'] == partial['reference_score'] == summary['new_score']
    assert quality['consistency']['status'] == consistency['status'] == summary['status']
    assert quality['consistency']['groups'][0]['total_scores'] == scores
    assert partial['model_answers'] == quality['discrimination']['model_scores'] == summary['new_model_scores']
    for root in [task, track]:
        for line in (root / 'records.jsonl').read_text().splitlines():
            record = json.loads(line)
            assert record['reference_score'] == summary['new_score'] and record['score_status'] == summary['status']
            for field in ['target', 'golden']:
                if isinstance(record.get(field), dict):
                    assert record[field]['reward'] == summary['new_score']
                    assert record[field]['status'] == summary['status']
                    assert record[field]['score_available'] == bool(scores)
                    assert record[field]['new_llm_calls'] == summary['new_llm_calls'] == 0
                    assert record[field]['imported_judge_runs'] == summary['valid_runs'] == 5
        metadata = read(root / 'metadata/task_metadata.json')
        assert metadata['reference_hybrid_score'] == summary['new_score']
        assert metadata['reference_fresh_valid_judge_runs'] == len(scores)
    archive = track / 'history' / summary['history']
    assert (archive / 'previous/golden/reward.json').is_file()
    previous = read(track / 'history/golden_judge_fix_summary.json')
    earlier = track / 'history' / previous['history'] / 'judge_runs'
    assert len(list(earlier.glob('golden_r[1-5]/reward-details.json'))) == 5
    costs = read(track / 'judge_costs.json')
    batch = costs['shared_account_batch']
    assert round(batch['usage_after_usd'] - batch['usage_before_usd'], 6) == batch['account_delta_usd'] == 17.31
    assert costs['total_cost_usd'] is summary['total_cost_usd'] is None
    assert costs['allocation'] == 'unallocated_shared_batch'
    assert len(costs['runs']) == summary['completed_runs']
    assert summary['new_llm_calls'] == 0 and summary['attempted_runs'] is None
    assert summary['representative'] == 'golden_r1' and group['run_ids'] == [f'golden_r{index}' for index in range(1, 6)]
    for row in costs['runs']:
        assert row['cost_usd'] is None and row['valid']
        assert (track / 'judge_runs' / row['run'] / 'reward-details.json').is_file()
    for folder in sorted((track / 'rollouts').iterdir()):
        candidate = folder / 'candidate.xlsx'
        result = engine.score(engine.extract(candidate, key), key)
        receipt = read(folder / 'reward.json')
        assert abs(result['score'] - receipt.get('programmatic_score', receipt['reward'])) < 0.000001, folder.name
        assert receipt['reward'] == summary['new_model_scores'][folder.name], folder.name
    blind = read(task / 'history/validation/judge_fixtures/saved_blind.json')
    assert finalizer.llm_evidence_errors({'reward': blind['agent']}, golden)
    print(json.dumps({'saved_valid_runs': len(scores), 'sample_sigma': sigma,
                      'current_score': summary['new_score'], 'requirement_met': requirement,
                      'invalid_attempts_excluded': summary['invalid_attempts_count']}))


def import_module(name, path):
    specification = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(module)
    return module


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--trajectory', type=Path)
    arguments = parser.parse_args()
    task = Path(__file__).resolve().parents[1]
    track = arguments.trajectory or task.parents[1] / 'ZCloud-电力监管-001-20261020-轨迹报告/powercor-2020-capex-reconciliation'
    counts = check_saved_cases(task)
    with tempfile.TemporaryDirectory() as temporary:
        scoring_main(temporary)
        pooling_main()
        access_main()
    check_saved_report(task, track)
    print(json.dumps(dict(status='PASS', **counts)))


if __name__ == '__main__':
    main()
