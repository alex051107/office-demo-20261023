"""Offline workbook extraction and deterministic atomic weighted scoring."""
from pathlib import Path
import argparse, ast, json, math, re, unicodedata
from openpyxl import load_workbook

def norm(x):
 s=unicodedata.normalize('NFKC',str(x or '')).casefold()
 return re.sub(r'\s+',' ',re.sub('[‐-―−]','-',s)).strip().rstrip('.;:')
def num(x):
 if isinstance(x,bool):return None
 if isinstance(x,(int,float)):return float(x) if math.isfinite(x) else None
 s=re.sub(r'[,$\s]','',str(x or '')).rstrip('%')
 if s.startswith('(') and s.endswith(')'):s='-'+s[1:-1]
 try:
  v=float(s);return v if math.isfinite(v) else None
 except ValueError:return None
def canon(x,aliases):
 n=norm(x)
 for key,values in aliases.items():
  if n in [norm(key)]+[norm(v) for v in values]:return norm(key)
 return n

def source_files(text, key):
 names=re.findall(r'[A-Za-z0-9_.-]+\.xlsx',str(text or ''),re.I)
 return names if names and all(name in key['source_locations'] for name in names) else []

def source_location(text, files, key):
 if not isinstance(text,str) or not files:return False
 named=re.findall(r'[A-Za-z0-9_.-]+\.xlsx',text,re.I)
 if any(name not in key['source_locations'] for name in named):return False
 locations=re.findall(r'\b(T\d{3}|Text)(?:!\$?([A-Z]+)\$?(\d+))?\b',text)
 if not locations:return False
 for sheet,column,row in locations:
  bounds=[key['source_locations'][name][sheet] for name in files if sheet in key['source_locations'][name]]
  if not bounds:return False
  if column and not any(1<=int(row)<=height and sum((ord(letter)-64)*26**power for power,letter in enumerate(reversed(column)))<=width for height,width in bounds):return False
 return True

def results_source(e,key,ident):
 notes=observations(e,key,['Results',ident,'source_note'])
 if len(notes)!=1:return False
 text=notes[0]['raw'];files=source_files(text,key)
 if not files and isinstance(text,str) and re.search(r'same.*(?:row\s*2|sources)',text,re.I):
  first=observations(e,key,['Results',key['interfaces']['Results']['expected_keys'][0],'source_note'])
  if first:text=first[0]['raw'];files=source_files(text,key)
 return source_location(text,files,key)

def stated_precision(observation,decimals,percent=False):
 raw=observation['raw'];number=num(raw)
 if number is None:return False
 number_format=observation.get('format','General').split(';')[0]
 cleaned=re.sub(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.', '', number_format)
 if '%' in cleaned and not percent:return False
 stored=number*100 if percent and '%' in cleaned and not isinstance(raw,str) else number
 if abs(stored-round(stored,decimals))<=1e-9:return True
 if re.search(r'[0#?]',cleaned):
  fraction=re.search(r'\.([0#?]+)',cleaned)
  return bool(fraction and len(fraction[1])==decimals and fraction[1].count('0')==decimals)
 return False

def extract(path,key):
 if not Path(path).is_file():return {'status':'OUTPUT_MISSING','sheets':{},'formulas':[]}
 try:
  w=load_workbook(path,data_only=True); fw=load_workbook(path,data_only=False)
 except Exception as e:return {'status':'JUDGE_ERROR','error':str(e),'sheets':{},'formulas':[]}
 sheets={}; formula_cells=[]; resolving=set()
 errors=[]
 ERR={'#DIV/0!','#VALUE!','#REF!','#NAME?','#N/A','#NUM!','#NULL!','#SPILL!','#CALC!'}
 for s_ in w:
  if s_.title not in key['interfaces']:continue
  for row_ in s_.iter_rows():
   for c_ in row_:
    if c_.data_type=='e' or (isinstance(c_.value,str) and c_.value.strip().upper() in ERR):errors.append(s_.title+'!'+c_.coordinate)
 def value(sh,coord):
  c=w[sh][coord];f=fw[sh][coord]
  if f.data_type!='f':return c.value
  formula_cells.append(sh+'!'+coord)
  if c.value is not None:return c.value
  ident=(sh,coord)
  if ident in resolving:return None
  resolving.add(ident)
  try:
   expr=f.value[1:].replace('$','')
   refs={}
   def ref(m):
    name=m.group(1) or m.group(2) or sh;cell=m.group(3);v=value(name,cell)
    if num(v) is None: raise ValueError('unresolved cell')
    token='v'+str(len(refs));refs[token]=num(v);return token
   expr=re.sub(r"(?:'([^']+)'!|([A-Za-z_][A-Za-z0-9_ ]*)!)?([A-Z]+[0-9]+)",ref,expr)
   def ev(n):
    if isinstance(n,ast.Expression):return ev(n.body)
    if isinstance(n,ast.Constant) and isinstance(n.value,(int,float)):return n.value
    if isinstance(n,ast.Name):return refs[n.id]
    if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)):return ev(n.operand)*(1 if isinstance(n.op,ast.UAdd) else -1)
    if isinstance(n,ast.BinOp):
     a,b=ev(n.left),ev(n.right)
     if isinstance(n.op,ast.Add):return a+b
     if isinstance(n.op,ast.Sub):return a-b
     if isinstance(n.op,ast.Mult):return a*b
     if isinstance(n.op,ast.Div):return a/b
    if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id.upper() in ['SUM','ROUND','ABS']:
     a=[ev(x) for x in n.args];return sum(a) if n.func.id.upper()=='SUM' else round(*a) if n.func.id.upper()=='ROUND' else abs(a[0])
    raise ValueError('unsupported formula')
   return ev(ast.parse(expr,mode='eval'))
  except Exception:return None
  finally:resolving.remove(ident)
 interface_errors=[];header_errors=[]
 for sheet,definition in key['interfaces'].items():
  if sheet not in w.sheetnames:
   interface_errors.append({'sheet':sheet,'reason':'required sheet missing or misnamed'})
   sheets[sheet]={'headers':[],'records':{},'extra':[],'row_keys':[],'actual_name':None};continue
  s=w[sheet]
  used_columns=[cell.column for row in fw[sheet] for cell in row if cell.value is not None]
  column_count=max(used_columns,default=0)
  raw_headers=[s.cell(1,column).value for column in range(1,column_count+1)]
  required_count=len(definition['headers'])
  required_headers=[header.strip() if isinstance(header,str) else header for header in raw_headers[:required_count]]
  if required_headers!=definition['headers'] or definition.get('exact_column_count',False) and column_count!=required_count:
   header_errors.append({'sheet':sheet,'reason':'required header prefix text/order mismatch after trimming','observed':raw_headers,'expected':definition['headers']})
  headers=[canon(header,definition.get('header_aliases',{})) for header in raw_headers[:required_count]]
  expected={tuple(canon(v,definition.get('value_aliases',{})) for v in x) for x in definition['expected_keys']}
  records={};extra=[];extra_business=[];row_keys=[];seen=set()
  for row in list(s)[1:]:
   record={h:{'raw':value(s.title,c.coordinate),'cell':s.title+'!'+c.coordinate,'format':c.number_format} for h,c in zip(headers,row) if h}
   if not any(cell['raw'] is not None and str(cell['raw']).strip() for cell in record.values()):continue
   ident=tuple(canon(record.get(h,{}).get('raw'),definition.get('value_aliases',{})) for h in definition['key'])
   row_keys.append([record.get(header,{}).get('raw') for header in definition['key']])
   records.setdefault('|'.join(ident),[]).append(record)
   if ident not in expected or ident in seen:extra.append(ident)
   business_field={'Results':'value','Analysis':'amount','Summary':'insurance_policy_benefits'}[sheet]
   business_value=record.get(business_field,{}).get('raw')
   if ident not in expected and business_value is not None and str(business_value).strip():extra_business.append({'row':row[0].row,'key':list(ident),'value':business_value})
   seen.add(ident)
  sheets[sheet]={'headers':headers,'raw_headers':raw_headers,'records':records,'extra':extra,'extra_business':extra_business,'row_keys':row_keys,'actual_name':s.title}
 w.close();fw.close()
 return {'status':'PARSED','sheets':sheets,'formulas':sorted(set(formula_cells)),'error_cells':errors,'interface_errors':interface_errors,'header_errors':header_errors}

def observations(e,key,ref):
 sheet,ident,field=ref; aliases=key['interfaces'][sheet].get('value_aliases',{})
 enc='|'.join(canon(x,aliases) for x in ident)
 return [r.get(field,{'raw':None,'cell':None,'format':''}) for r in e['sheets'].get(sheet,{}).get('records',{}).get(enc,[])]
def amount(e,key,ref,percent=False):
 oo=observations(e,key,ref);vv=[]
 for o in oo:
  v=num(o['raw'])
  if v is None:return None
  if percent and '%' in o['format'] and not isinstance(o['raw'],str):v*=100
  if ref[0]=='Analysis' and ref[2]=='amount':
   units=observations(e,key,[ref[0],ref[1],'unit'])
   u=norm(units[0]['raw']) if units else ''
   if key.get('amount_base_unit')=='usd millions':
    if u in ['usd','usd dollars','dollars']:v/=1e6
    elif u in ['usd thousands','usd thousand','thousands of usd']:v/=1000
   else:
    if u in ['usd','usd dollars','dollars']:v/=1000
    elif u in ['usd millions','usd million','millions of usd']:v*=1000
  vv.append(v)
 if not vv or max(vv)-min(vv)>1e-8:return None
 return vv[0]
def score(e,key):
 verdicts={}; criteria=key['rows']+key['tables']+key['checks']
 for c in key['rows']+key['tables']:
  if c in key['rows']:
   ref=['Results',[c['label']],'value'];spec=c;gold=c['gold']
  else:
   field,spec=next(iter(c['value_columns'].items()));r=c['gold'][0]
   ref=[c['sheet'],[r[k] for k in c['key']],field];gold=r[field]
  v=amount(e,key,ref,spec.get('percent_forms',False));met=v is not None and abs(v-gold)<=spec.get('abs',0)+spec.get('rel',0)*abs(gold)+1e-9
  verdicts[c['criterion_id']]={'credit':float(met),'observed':v,'expected':gold,'evidence':observations(e,key,ref)}
 for c in key['checks']:
  typ=c['kind']; observed=None;credit=0.
  if typ=='aggregate':
   pf=c.get('percent_refs') or [False]*len(c['refs'])
   vals=[amount(e,key,r,p_) for r,p_ in zip(c['refs'],pf)]; observed=vals
   credit=float(all(v is not None for v in vals) and abs(sum(a*b for a,b in zip(vals,c['coefficients'])))<=c['abs']+1e-8)
  elif typ=='source':
   oo=observations(e,key,c['ref']);observed=oo
   credit=float(bool(oo) and all(source_files(o['raw'],key) for o in oo))
  elif typ=='header':
   d=e['sheets'].get(c['sheet'],{});observed=d.get('headers',[])
   nonempty=any(x.get('raw') not in [None,''] for rr in d.get('records',{}).values() for r in rr for h,x in r.items() if h not in key['interfaces'][c['sheet']]['key'])
   credit=float(d.get('actual_name')==c['sheet'] and not any(error['sheet']==c['sheet'] for error in e.get('header_errors',[])))
  elif typ=='error_cells':
   observed=e.get('error_cells',[]);credit=float(bool(observed))
  elif typ=='unsupported_records':
   extra=[x for s in e['sheets'].values() for x in s.get('extra',[])];observed=extra
   credit=float(bool(extra))
  elif typ=='required_sources':
   observed=[observations(e,key,['Results',ident,'source_note']) for ident in key['interfaces']['Results']['expected_keys']]
   observed.extend(observations(e,key,['Analysis',ident,field]) for ident in key['interfaces']['Analysis']['expected_keys'] for field in ['source_document','source_page'])
   credit=float(any(len(notes)!=1 or not isinstance(notes[0]['raw'],str) or not notes[0]['raw'].strip() or num(notes[0]['raw']) is not None for notes in observed))
  elif typ=='results_locations':
   observed=[results_source(e,key,ident) for ident in key['interfaces']['Results']['expected_keys']]
   credit=float(not all(observed))
  elif typ=='analysis_locations':
   observed=[]
   for ident in key['interfaces']['Analysis']['expected_keys']:
    documents=observations(e,key,['Analysis',ident,'source_document']);locations=observations(e,key,['Analysis',ident,'source_page'])
    observed.append(len(documents)==len(locations)==1 and source_location(locations[0]['raw'],source_files(documents[0]['raw'],key),key))
   credit=float(not all(observed))
  elif typ=='row_order':
   observed={sheet:definition.get('row_keys',[]) for sheet,definition in e['sheets'].items()}
   credit=float(any([ident for ident in observed.get(sheet,[]) if ident in definition['expected_keys']]!=definition['expected_keys'] for sheet,definition in key['interfaces'].items()))
  elif typ=='units':
   observed=[]
   for sheet in ['Results','Analysis']:
    for index,ident in enumerate(key['interfaces'][sheet]['expected_keys']):
     expected_unit=key['rows'][index]['unit'] if sheet=='Results' else 'USD millions'
     units=observations(e,key,[sheet,ident,'unit']);observed.append(len(units)==1 and units[0]['raw']==expected_unit)
   credit=float(not all(observed))
  elif typ=='precision':
   observed=observations(e,key,c['ref'])
   credit=float(len(observed)!=1 or not stated_precision(observed[0],c['decimals'],c.get('percent',False)))
  else:raise ValueError(typ)
  verdicts[c['criterion_id']]={'credit':credit,'evidence':observed}
 gate={'failed':bool(e.get('interface_errors')),'errors':e.get('interface_errors',[])}
 if gate['failed']:
  for criterion in key['rows']+key['tables']:
   if criterion['weight']>0:verdicts[criterion['criterion_id']]['credit']=0.0
 total=sum(c['weight'] for c in criteria if c['weight']>0)
 earned=sum(c['weight']*verdicts[c['criterion_id']]['credit'] for c in criteria)
 for c in criteria:verdicts[c['criterion_id']].update(weight=c['weight'],dimension=c['dimension'])
 bygroup={}
 for c in criteria:
  bygroup.setdefault(c['budget_group'],0);bygroup[c['budget_group']]+=c['weight']*verdicts[c['criterion_id']]['credit']
 positive=[c for c in criteria if c['weight']>0]
 baseline=sum(c['baseline_weight']*verdicts[c['criterion_id']]['credit'] for c in key['rows'])/sum(c['baseline_weight'] for c in key['rows'])
 result={'status':'SCORED','task_id':key['task_id'],'score':min(1,max(0,earned/total)),'total_weight':total,'earned':earned,'groups':bygroup,
  'baseline_roots_score':baseline,'equal_criterion_score':sum(verdicts[c['criterion_id']]['credit'] for c in positive)/len(positive),
  'equal_business_score':sum(verdicts[c['criterion_id']]['credit'] for c in key['rows']+key['tables'])/(len(key['rows'])+len(key['tables'])),
  'verdicts':verdicts,'core_cap':False,'llm_judge_fraction':0,'interface_gate':gate}
 if e['status']=='JUDGE_ERROR':result.update(status='JUDGE_ERROR',score=None,error=e.get('error'))
 if e['status']=='OUTPUT_MISSING':result.update(status='OUTPUT_MISSING',score=0.0)
 reasons=[]
 if verdicts['S01']['credit']:reasons.append('a required source/support cell is empty or numeric-only')
 if verdicts['U01']['credit']:reasons.append('a required unit differs from the instruction')
 for sheet,definition in e['sheets'].items():
  for record in definition.get('extra_business',[]):reasons.append({'reason':'fabricated extra business record','sheet':sheet,**record})
 result['post_pool_rules']=[{'name':'required_header_cap','triggered':bool(e.get('header_errors')),'cap':0.69,'reasons':e.get('header_errors',[])},
  {'name':'instruction_must_clause_cap','triggered':bool(reasons),'cap':0.69,'reasons':reasons}]
 result['score_before_post_pool_rules']=result['score']
 for rule in result['post_pool_rules']:
  if rule['triggered']:result['score']=min(result['score'],rule['cap'])
 if gate['failed']:result['score']=0.0
 return result

def main():
 p=argparse.ArgumentParser();p.add_argument('--key',type=Path,required=True);p.add_argument('--candidate',type=Path,required=True);p.add_argument('--out',type=Path,required=True)
 a=p.parse_args();key=json.loads(a.key.read_text());path=a.candidate if a.candidate.is_file() else a.candidate/'answer.xlsx'
 e=extract(path,key);r=score(e,key);a.out.mkdir(parents=True,exist_ok=True)
 for f,d in [('extraction.json',e),('score.json',r)]: (a.out/f).write_text(json.dumps(d,ensure_ascii=False,indent=2,default=str)+'\n')
 (a.out/'reward.txt').write_text(str(r['score'] if r['score'] is not None else '')+'\n');print(json.dumps({'status':r['status'],'score':r['score']}))
 return 2 if r['status']=='JUDGE_ERROR' else 0
if __name__=='__main__':raise SystemExit(main())
