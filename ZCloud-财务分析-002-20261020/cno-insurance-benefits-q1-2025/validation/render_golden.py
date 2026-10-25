import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--office', default=shutil.which('soffice'))
    arguments = parser.parse_args()
    if not arguments.office:
        parser.error('Specify the LibreOffice executable with --office.')
    task = Path(__file__).resolve().parents[1]
    root = task.parents[1]
    golden = root / (task.parent.name + '-轨迹报告') / task.name / 'golden'
    candidate = golden / 'candidate.xlsx'
    renders = golden / 'renders'
    receipt = []
    workbook = openpyxl.load_workbook(candidate)
    names = workbook.sheetnames
    workbook.close()
    with tempfile.TemporaryDirectory(prefix='render_', dir=task / 'validation') as temporary:
        staging = Path(temporary)
        for index, name in enumerate(names, 1):
            copy = openpyxl.load_workbook(candidate)
            for other in list(copy):
                if other.title != name:
                    copy.remove(other)
            sheet = copy[name]
            sheet.sheet_state = 'visible'
            sheet.sheet_view.tabSelected = True
            copy.active = 0
            for view in copy.views:
                view.activeTab = 0
                view.firstSheet = 0
            sheet.print_area = f'A1:{openpyxl.utils.get_column_letter(sheet.max_column)}{sheet.max_row}'
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1
            sheet.page_setup.orientation = 'landscape'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            label = f'sheet-{index:02d}'
            source = staging / (label + '.xlsx')
            copy.save(source)
            copy.close()
            conversion = subprocess.run(
                [arguments.office, '-env:UserInstallation=' + (staging / 'profile').as_uri(),
                 '--headless', '--convert-to', 'pdf:calc_pdf_Export', '--outdir', str(staging), str(source)],
                capture_output=True, text=True, timeout=60,
            )
            if conversion.returncode:
                raise RuntimeError(f'{name}: LibreOffice export failed')
            pdf = renders / (label + '.pdf')
            shutil.copyfile(staging / pdf.name, pdf)
            subprocess.run(['pdftoppm', '-r', '110', '-png', '-singlefile', str(pdf), str(renders / label)],
                           check=True, capture_output=True, timeout=60)
            text = subprocess.run(['pdftotext', '-layout', str(pdf), '-'],
                                  check=True, capture_output=True, text=True, timeout=30).stdout
            expected = {'Results': 'source_note', 'Analysis': 'source_document',
                        'Summary': 'insurance_policy_benefits'}[name]
            if expected not in text:
                raise AssertionError(f'{name}: PDF does not contain its own header')
            info = subprocess.run(['pdfinfo', str(pdf)], check=True, capture_output=True,
                                  text=True, timeout=30).stdout
            if not any(line.split() == ['Pages:', '1'] for line in info.splitlines()):
                raise AssertionError(f'{name}: expected one PDF page')
            receipt.append({'sheet': name, 'pdf': 'renders/' + pdf.name,
                            'png': 'renders/' + label + '.png', 'page_count': 1,
                            'header_verified': expected, 'pdf_sha256': hashlib.sha256(pdf.read_bytes()).hexdigest()})
    metadata = json.loads((golden / 'visual_review.json').read_text())
    metadata.update({'candidate_sha256': hashlib.sha256(candidate.read_bytes()).hexdigest(),
                     'render_status': 'COMPLETED', 'visual_inspection_status': 'PENDING',
                     'scope': '临时渲染副本仅保留目标工作表，原参考答案不变；LibreOffice导出A3横向单页PDF，Poppler生成110dpi PNG，并核对每份PDF的专属表头。',
                     'visual_findings': [], 'inspection_coverage': [],
                     'inspection_method': '等待逐张查看本次生成的三张PNG。'})
    metadata.pop('layout_finding_handling', None)
    (golden / 'visual_review.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + '\n')
    (task / 'validation/render_receipt.json').write_text(json.dumps(
        {'renderer': metadata['renderer'], 'candidate_sha256': metadata['candidate_sha256'],
         'sheets': receipt}, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'rendered_sheets': names, 'distinct_pdf_count': len({row['pdf_sha256'] for row in receipt})}))


if __name__ == '__main__':
    main()
