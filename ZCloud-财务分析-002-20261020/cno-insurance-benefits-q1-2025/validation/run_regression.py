"""Read-only offline regression entry for the task and its saved report."""
import argparse
import hashlib
import importlib.util
import json
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True

"""Offline interface and instruction-constraint regression workbooks."""
import argparse
import copy
import importlib.util
import json
import sys
import tempfile
from pathlib import Path
from openpyxl import load_workbook
sys.dont_write_bytecode = True
scoring_TASK = Path(__file__).resolve().parents[1]

def scoring_load_module(name, path):
    specification = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(module)
    return module

def scoring_reverse_columns(worksheet):
    values = [[copy.copy(cell) for cell in row] for row in worksheet]
    for row_number, row in enumerate(values, 1):
        for column_number, cell in enumerate(reversed(row), 1):
            destination = worksheet.cell(row_number, column_number)
            destination.value = cell.value
            destination.number_format = cell.number_format

def scoring_swap_rows(worksheet):
    rows = [[copy.copy(cell) for cell in worksheet[row]] for row in (2, 3)]
    for row_number, row in zip((3, 2), rows):
        for column_number, cell in enumerate(row, 1):
            worksheet.cell(row_number, column_number, cell.value)
            worksheet.cell(row_number, column_number).number_format = cell.number_format

def scoring_mutate(workbook, case):
    worksheet = workbook[case.get('sheet', 'Results')]
    mutation = case['mutation']
    if mutation == 'review_wrong_names_reverse_all':
        workbook['Results'].title = 'WrongSheet1'
        workbook['Analysis'].title = 'WrongSheet2'
        for current in workbook:
            scoring_reverse_columns(current)
    elif mutation == 'rename':
        worksheet.title = 'Wrong_' + worksheet.title
    elif mutation == 'reverse_columns':
        scoring_reverse_columns(worksheet)
    elif mutation == 'extra_header':
        worksheet.cell(1, worksheet.max_column + 1, 'extra')
    elif mutation == 'missing_header':
        worksheet.cell(1, worksheet.max_column).value = None
    elif mutation == 'shift_header':
        worksheet.insert_rows(1)
    elif mutation == 'header_case':
        worksheet['A1'] = worksheet['A1'].value.upper()
    elif mutation == 'set_cell':
        worksheet[case['cell']] = case['value']
        if 'format' in case:
            worksheet[case['cell']].number_format = case['format']
    elif mutation == 'swap_rows':
        scoring_swap_rows(worksheet)
    elif mutation in {'extra_record', 'duplicate_record', 'blank_key_record'}:
        values = [cell.value for cell in worksheet[2]]
        if mutation == 'extra_record':
            values[0] = 'Unexpected business record'
        if mutation == 'blank_key_record':
            values[0] = None
        worksheet.append(values)
    elif mutation == 'missing_record':
        worksheet.delete_rows(2)
    else:
        raise ValueError(mutation)

def scoring_definitions():
    cases = [{'id': 'review_wrong_structure', 'mutation': 'review_wrong_names_reverse_all', 'gate': True, 'old_score_from_review': 1.0}]
    for sheet in ('Results', 'Analysis', 'Summary'):
        for mutation in ('rename', 'reverse_columns', 'missing_header', 'shift_header', 'header_case'):
            cases.append({'id': sheet.lower() + '_' + mutation, 'sheet': sheet, 'mutation': mutation, 'gate': mutation == 'rename', 'cap': mutation != 'rename'})
        for mutation in ('swap_rows', 'extra_record', 'duplicate_record', 'blank_key_record', 'missing_record'):
            cases.append({'id': sheet.lower() + '_' + mutation, 'sheet': sheet, 'mutation': mutation, 'cap': mutation in {'extra_record', 'blank_key_record'}, 'penalty': 'O01' if mutation in {'swap_rows', 'missing_record'} else 'X2'})
    for label, cell, value, penalty in [('results_empty_source_d2', 'D2', None, 'S01'), ('results_empty_source_d3', 'D3', ' ', 'S01'), ('results_numeric_source', 'D2', 123, 'S01'), ('results_generic_source', 'D2', 'CNO filings', 'S02'), ('results_fake_file', 'D2', 'Imaginary.xlsx T010!G19', 'S02'), ('results_fake_table', 'D2', 'CNO_2025-03-31_10Q.xlsx T999!G19', 'S02'), ('results_fake_cell', 'D2', 'CNO_2025-03-31_10Q.xlsx T010!ZZ99999', 'S02'), ('results_wrong_unit', 'C2', 'USD thousands', 'U01'), ('results_unit_case', 'C3', 'Percent', 'U01'), ('results_percentage_general', 'B3', 12.13850088530396, 'DP2'), ('results_percentage_unrounded_coarse_format', 'B3', 12.13850088530396, 'DP2'), ('results_amount_extra_decimal', 'B2', 61.71, 'DP1'), ('results_percentage_outside_tolerance', 'B3', 12.1441, 'R02')]:
        case = {'id': label, 'mutation': 'set_cell', 'cell': cell, 'value': value, 'penalty': penalty, 'cap': penalty in {'S01', 'U01'}}
        if cell.startswith('B'):
            case['format'] = '0.0' if 'coarse_format' in label else '0.00' if 'tolerance' in label else 'General'
        cases.append(case)
    for label, cell, value, penalty in [('analysis_fake_file', 'E2', 'Imaginary.xlsx', 'SD1'), ('analysis_empty_location', 'F3', None, 'SP1'), ('analysis_fake_table', 'F2', 'T999', 'SP1'), ('analysis_wrong_unit', 'D2', 'USD thousand', 'U01')]:
        cases.append({'id': label, 'sheet': 'Analysis', 'mutation': 'set_cell', 'cell': cell, 'value': value, 'penalty': penalty, 'cap': label in {'analysis_empty_location', 'analysis_wrong_unit'}})
    for cell in ('E2', 'E3', 'F2', 'F3'):
        for label, value in (('empty', None), ('numeric', 123), ('numeric_text', '123')):
            cases.append({'id': 'analysis_' + cell.lower() + '_' + label, 'sheet': 'Analysis', 'mutation': 'set_cell', 'cell': cell, 'value': value, 'penalty': 'S01', 'cap': True})
    return cases

def scoring_verify_finalize(finalizer, result, weights):
    criteria = [{'name': name, 'weight': abs(weights[name]), 'value': 1 - verdict['credit'] if weights[name] < 0 else verdict['credit'], 'negate': weights[name] < 0} for name, verdict in result['verdicts'].items()]
    details = {'reward': [{'kind': 'programmatic', 'criteria': criteria}], 'interface_gate': result['interface_gate'], 'post_pool_rules': result['post_pool_rules']}
    with tempfile.TemporaryDirectory() as temporary:
        graded = Path(temporary) / 'reward.json'
        graded.write_text(json.dumps({'programmatic': result['score']}))
        graded.with_name('reward-details.json').write_text(json.dumps(details))
        verdict = finalizer.compute(argparse.Namespace(graded=str(graded), graded_rc=0, candidate=scoring_TASK / 'solution/golden_output/answer.xlsx'))
        assert verdict['reward'] == result['score'] and verdict['verifier_error'] == 0
        details['reward'].append({'kind': 'agent', 'criteria': [{'id': 'L01', 'value': 1.0, 'raw': 'no', 'weight': 3, 'negate': True, 'reasoning': '无法解析候选文件，文件无法解析，无法检查来源。'}]})
        graded.with_name('reward-details.json').write_text(json.dumps(details))
        unavailable = finalizer.compute(argparse.Namespace(graded=str(graded), graded_rc=0, candidate=scoring_TASK / 'solution/golden_output/answer.xlsx'))
        assert unavailable['reward'] == 0 and unavailable['verifier_error'] == 1
        assert unavailable['llm_evidence_errors']

def scoring_main(temporary):
    parser = argparse.ArgumentParser()
    parser.add_argument('--baseline', action='store_true')
    arguments = argparse.Namespace(baseline=False)
    engine = scoring_load_module('regression_engine', scoring_TASK / 'tests/score_output.py')
    finalizer = scoring_load_module('regression_finalizer', scoring_TASK / 'tests/finalize.py')
    key = json.loads((scoring_TASK / 'tests/results_key.json').read_text())
    weights = {item['criterion_id']: item['weight'] for item in key['rows'] + key['tables'] + key['checks']}
    directory = Path(temporary) / 'negative_cases'
    directory.mkdir(parents=True, exist_ok=True)
    metadata = directory / 'cases.json'
    previous_metadata = json.loads(metadata.read_text()) if metadata.exists() else {}
    previous = {case['id']: case for case in previous_metadata.get('cases', []) + previous_metadata.get('controls', [])}
    golden = scoring_TASK / 'solution/golden/answer.xlsx'
    baseline = engine.score(engine.extract(golden, key), key)
    assert baseline['score'] == 1.0, baseline
    results = []
    for definition in scoring_definitions():
        workbook = load_workbook(golden)
        scoring_mutate(workbook, definition)
        path = directory / (definition['id'] + '.xlsx')
        workbook.save(path)
        workbook.close()
        result = engine.score(engine.extract(path, key), key)
        scoring_verify_finalize(finalizer, result, weights)
        pooled, counted, broken = finalizer.pooled_score({'reward': {'criteria': [{'weight': abs(weights[name]), 'value': 1 - verdict['credit'] if weights[name] < 0 else verdict['credit'], 'negate': weights[name] < 0} for name, verdict in result['verdicts'].items()]}})
        pooled = finalizer.apply_post_pool_rules(result, pooled)
        assert not broken and counted == len(weights) and (abs(pooled - result['score']) < 1e-06)
        failed = [name for name, verdict in result['verdicts'].items() if verdict['credit'] == (1 if weights[name] < 0 else 0)]
        record = dict(definition, workbook=path.name, expected_outcome='gate 0' if definition.get('gate') else 'cap <=0.69' if definition.get('cap') else 'penalty <1', old_score_from_review=definition.get('old_score_from_review'), old_program_score=previous.get(definition['id'], {}).get('old_program_score'), round1_program_score=previous.get(definition['id'], {}).get('round1_program_score', previous.get(definition['id'], {}).get('new_program_score')), new_program_score=None if arguments.baseline else result['score'], triggered=failed, interface_gate=result.get('interface_gate'), post_pool_rules=result.get('post_pool_rules'))
        if not arguments.baseline:
            assert result['score'] < 1.0, record
            if definition.get('gate'):
                assert result['score'] == 0 and result['interface_gate']['failed'], record
            elif definition.get('cap'):
                assert result['score'] <= 0.69 and any((rule['triggered'] for rule in result['post_pool_rules'])), record
            if definition.get('penalty'):
                assert definition['penalty'] in failed, record
        results.append(record)
    controls = []
    if not arguments.baseline:
        for label in ('support_sheet_unconstrained', 'rounded_stored_values', 'percentage_excel_format', 'styling_only_extra_column', 'trailing_header_spaces', 'leading_header_spaces', 'results_percentage_wrong_format', 'formatted_unrounded_value', 'results_extra_header', 'analysis_extra_header', 'summary_extra_header', 'trailing_duplicate_header'):
            workbook = load_workbook(golden)
            if label == 'support_sheet_unconstrained':
                support = workbook.create_sheet('Support')
                support.append(['Anything', '#DIV/0!', 123.4567])
            elif label == 'rounded_stored_values':
                workbook['Results']['B2'].number_format = 'General'
                workbook['Results']['B3'] = 12.14
                workbook['Results']['B3'].number_format = 'General'
            elif label == 'percentage_excel_format':
                workbook['Results']['B3'] = 61.7 / 508.3
                workbook['Results']['B3'].number_format = '0.00%'
            elif label == 'styling_only_extra_column':
                workbook['Results']['E1'].number_format = '0.00'
            elif label in {'trailing_header_spaces', 'leading_header_spaces'}:
                for worksheet in workbook:
                    for cell in worksheet[1]:
                        cell.value = cell.value + '  ' if label == 'trailing_header_spaces' else '  ' + cell.value
            elif label == 'results_percentage_wrong_format':
                workbook['Results']['B3'] = 12.14
                workbook['Results']['B3'].number_format = '0.0'
            elif label == 'formatted_unrounded_value':
                workbook['Results']['B2'] = 61.71
                workbook['Results']['B2'].number_format = '0.0'
            elif label == 'trailing_duplicate_header':
                for worksheet, header in [(workbook['Results'], 'value'), (workbook['Analysis'], 'amount'), (workbook['Summary'], 'insurance_policy_benefits')]:
                    extra_column = worksheet.max_column + 1
                    worksheet.cell(1, extra_column, header)
                    worksheet.cell(2, extra_column, 123)
                    worksheet.cell(3, extra_column, 456)
            else:
                worksheet = workbook[label.split('_')[0].title()]
                worksheet.cell(1, worksheet.max_column + 1, 'extra')
                worksheet.cell(2, worksheet.max_column, 'Additional permitted support')
            path = directory / (label + '.xlsx')
            workbook.save(path)
            workbook.close()
            score = engine.score(engine.extract(path, key), key)['score']
            assert score == 1.0, (label, score)
            controls.append({'id': label, 'workbook': path.name, 'expected_program_score': 1.0, 'program_score': score, 'round1_program_score': previous.get(label, {}).get('round1_program_score', previous.get(label, {}).get('new_program_score'))})
    metadata.write_text(json.dumps({'scope': 'deterministic programmatic subset; no new text judgment', 'precision_rule': 'p decimals satisfied by number_format OR stored value equal to rounding to p decimals; percentage formats use percentage points for stored comparison', 'old_score_scope': 'review value only where present; other old_program_score values are offline v1 replays on the formatted golden', 'golden_program_score': baseline['score'], 'cases': results, 'controls': controls}, ensure_ascii=False, indent=2) + '\n')
    scoring_verify_finalize(finalizer, baseline, weights)
    print(json.dumps({'negative_cases': len(results), 'controls': len(controls), 'golden': baseline['score'], 'finalize_policy_and_unreadable_judge_checks': len(results) + 1, 'mode': 'baseline' if arguments.baseline else 'verified'}))

"""Check accepted-run statistics, unavailable scores, and historical exclusion offline."""
import argparse
import hashlib
import importlib.util
import json
import statistics
from pathlib import Path
saved_TASK = Path(__file__).resolve().parents[1]

def saved_read(path):
    return json.loads(path.read_text())

def saved_load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def saved_main(track):
    parser = argparse.ArgumentParser()
    parser.add_argument('--trajectory', type=Path, required=True)
    track = Path(track)
    finalizer = saved_load('golden_record_finalizer', saved_TASK / 'tests/finalize.py')
    engine = saved_load('golden_record_engine', saved_TASK / 'tests/score_output.py')
    key = saved_read(saved_TASK / 'tests/results_key.json')
    candidate = track / 'golden/candidate.xlsx'
    result = engine.score(engine.extract(candidate, key), key)
    assert result['score'] == 1.0 and (not result['interface_gate']['failed'])
    assert not any((rule['triggered'] for rule in result['post_pool_rules']))
    summary = saved_read(track / 'history/golden_judge_fix_summary.json')
    assert summary == saved_read(saved_TASK / 'history/validation/golden_judge_fix_summary.json')
    consistency = saved_read(track / 'consistency.json')
    group = consistency['groups'][0]
    judges = []
    for identifier in group['run_ids']:
        folder = track / 'judge_runs' / identifier
        reward = saved_read(folder / 'reward.json')
        details = saved_read(folder / 'reward-details.json')
        judge = saved_read(folder / 'judge.json')
        assert reward['verifier_error'] == judge['verifier_error'] == 0
        assert finalizer.llm_evidence_errors(details, candidate) == []
        pooled, counted, broken = finalizer.pooled_score(details)
        score = round(finalizer.apply_post_pool_rules(details, pooled), 6)
        assert not broken and counted == len(saved_read(saved_TASK / 'rubrics.json')['items'])
        assert reward['reward'] == judge['score_final'] == score
        assert judge['reward_details_sha256'] == hashlib.sha256((folder / 'reward-details.json').read_bytes()).hexdigest()
        judges.append(judge)
    scores = [judge['score_final'] for judge in judges]
    expected_sigma = statistics.stdev(scores) if len(scores) > 1 else None
    assert group['n'] == summary['valid_runs'] == len(scores)
    assert group['total_score']['values'] == scores
    assert group['total_score']['stddev_sample'] == expected_sigma == summary['total_sigma_sample']
    assert len(group['criteria']) == len(saved_read(saved_TASK / 'rubrics.json')['items'])
    for criterion in group['criteria']:
        values = [next((item['score'] for item in judge['criteria'] if item['criterion_id'] == criterion['criterion_id'])) for judge in judges]
        sigma = statistics.stdev(values) if len(values) > 1 else None
        assert criterion['values'] == values and criterion['n'] == len(values)
        assert criterion['stddev_sample'] == sigma
    assert group['requirement_met'] == (len(scores) == 5 and expected_sigma is not None and (expected_sigma < 0.05) and all((item['pass_stddev_lt_0_05'] for item in group['criteria'])))
    quality = saved_read(track / 'quality_report.json')
    partial = saved_read(track / 'discrimination_cases.json')
    new_score = summary['new_score']
    assert quality['target']['score_final'] == quality['discrimination']['reference_score'] == partial['reference_score'] == new_score
    assert quality['consistency']['status'] == consistency['status'] == summary['status']
    assert quality['consistency']['groups'][0]['total_scores'] == scores
    assert partial['model_answers'] == quality['discrimination']['model_scores'] == summary['new_model_scores']
    for root in [saved_TASK, track]:
        for line in (root / 'records.jsonl').read_text().splitlines():
            record = json.loads(line)
            assert record['reference_score'] == new_score
            assert record['score_status'] == summary['status']
        metadata = saved_read(root / 'metadata/task_metadata.json')
        assert metadata['reference_hybrid_score'] == new_score
        assert metadata['reference_fresh_valid_judge_runs'] == len(scores)
    if not scores:
        reward = saved_read(track / 'golden/reward.json')
        judge = saved_read(track / 'golden/judge.json')
        assert reward['verifier_error'] == 1 and (not reward['score_available'])
        assert judge['score_final'] is None and (not judge['judge_completed'])
        assert not list((track / 'judge_runs').glob('golden_r[1-5]'))
    history = track / 'history' / summary['history']
    assert (history / 'golden/reward.json').is_file()
    archived = sorted((history / 'judge_runs').glob('golden_r[1-5]/reward-details.json'))
    assert len(archived) == 5
    blind = saved_read(saved_TASK / 'history/validation/judge_fixtures/saved_blind.json')
    assert finalizer.llm_evidence_errors({'reward': blind['agent']}, candidate)
    for label, score in summary['new_model_scores'].items():
        assert saved_read(track / 'rollouts' / label / 'reward.json')['reward'] == score
    print(json.dumps({'saved_valid_runs': len(scores), 'sample_sigma': expected_sigma, 'archived_runs': len(archived), 'programmatic_golden': 1.0, 'current_score': new_score, 'requirement_met': group['requirement_met']}))

"""Regression checks for candidate-access evidence and saved judgments."""
import argparse
import copy
import importlib.util
import json
import subprocess
import sys
import tempfile
from pathlib import Path
access_TASK = Path(__file__).resolve().parents[1]

def access_load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

def access_main():
    finalizer = access_load('artifact_finalizer', access_TASK / 'tests/finalize.py')
    checker = access_load('artifact_checker', access_TASK / 'tests/llm_artifact_check.py')
    golden = access_TASK / 'solution/golden_output/answer.xlsx'
    blind = json.loads((access_TASK / 'history/validation/judge_fixtures/saved_blind.json').read_text())
    with tempfile.TemporaryDirectory() as temporary:
        folder = Path(temporary)
        graded = folder / 'graded/reward.json'
        graded.parent.mkdir()
        graded.write_text('{"programmatic": 1.0}')
        valid = {'reward': [{'kind': 'programmatic', 'criteria': [{'name': 'control', 'weight': 10, 'value': 1.0, 'negate': False}]}, {'kind': 'agent', 'criteria': [{'id': name, 'weight': 3, 'value': 1.0, 'negate': True, 'reasoning': 'Results!D2 contains a source citation with no unsupported assertion.'} for name in ['L01', 'L02', 'L03']]}]}
        cases = [('evidence_backed_control', valid, False), ('saved_blind_C_golden_r4', {'reward': [valid['reward'][0], blind['agent']]}, True)]
        for name, reason, expected in [('no_cell', 'No unsupported assertions were found.', True), ('source_only_cell', 'The source Text!A131 supports this conclusion.', True), ('nonexistent_candidate_cell', 'Results!ZZ9999 supports this conclusion.', True), ('missing_with_cell', 'Results!D2 cannot be read because candidate workbook is missing.', True), ('unreadable_with_cell', 'Results!D2：无法解析候选文件。', True), ('missing_source_not_artifact', '候选 Results!D2 引用了不存在的来源文件。', False), ('missing_source_english', 'The candidate workbook Results!D2 cites a missing source file.', False), ('quoted_cell', "'Results'!$D$2 contains no audit-completion claim.", False)]:
            detail = copy.deepcopy(valid)
            detail['reward'][1]['criteria'][0]['reasoning'] = reason
            cases.append((name, detail, expected))
        for name, detail, expected_error in cases:
            graded.with_name('reward-details.json').write_text(json.dumps(detail))
            output = folder / name / 'reward.json'
            subprocess.run([sys.executable, str(access_TASK / 'tests/finalize.py'), '--graded', str(graded), '--graded-rc', '0', '--candidate', str(golden), '--out', str(output)], check=True)
            receipt = json.loads(output.read_text())
            assert bool(receipt['verifier_error']) == expected_error, (name, receipt)
            if expected_error:
                assert receipt['reward'] == 0
                exit_message = json.loads(output.with_name('reward_exit_message.json').read_text())
                assert exit_message['exit_code'] == 'judge:artifact_access_error', name
            else:
                assert receipt['reward'] == 1.0, name
                assert not output.with_name('reward_exit_message.json').exists(), name
        saved_valid = 0
        for fixture in (access_TASK / 'history/validation/judge_fixtures').glob('saved_valid_*.json'):
            payload = json.loads(fixture.read_text())
            assert checker.evidence_errors({'reward': payload['agent']}, golden) == [], fixture.name
            saved_valid += 1
        assert checker.evidence_errors({'reward': blind['agent']}, folder / 'missing.xlsx') == []
    print(json.dumps({'artifact_access_regressions': len(cases), 'saved_blind_cases': 1, 'saved_valid_cases': saved_valid, 'missing_candidate_keeps_program_gate': True}))


def check_saved_cases(task):
    engine = import_module('case_inventory_engine', task / 'tests/score_output.py')
    key = json.loads((task / 'tests/results_key.json').read_text())
    counts = {}
    for category in ('negative_cases', 'positive_cases'):
        folder = task / 'validation' / category
        cases = json.loads((folder / 'cases.json').read_text())['cases']
        assert {entry['file'] for entry in cases} == {path.name for path in folder.glob('*.xlsx')}
        for entry in cases:
            result = engine.score(engine.extract(folder / entry['file'], key), key)
            expected = entry.get('new_program_score', entry.get('program_score', entry.get('expected_program_score')))
            assert expected is not None and abs(result['score'] - expected) < 0.000001, entry['file']
            assert result['score'] == 1.0 if category == 'positive_cases' else result['score'] < 1.0
        counts[category] = len(cases)
    return counts


def check_saved_report(task, track):
    engine = import_module('saved_report_engine', task / 'tests/score_output.py')
    key = json.loads((task / 'tests/results_key.json').read_text())
    golden = track / 'golden/candidate.xlsx'
    assert golden.read_bytes() == (task / 'solution/golden_output/answer.xlsx').read_bytes()
    assert engine.score(engine.extract(golden, key), key)['score'] == 1.0
    reward = json.loads((track / 'golden/reward.json').read_text())
    judge = json.loads((track / 'golden/judge.json').read_text())
    consistency = json.loads((track / 'consistency.json').read_text())
    assert reward['verifier_error'] == 1 and reward['score_available'] is False
    assert judge['score_final'] is None and judge['judge_completed'] is False
    assert consistency['status'] == 'NO_VALID_NEW_JUDGMENT'
    assert all(group['n'] == 0 and not group['run_ids'] and not group['requirement_met']
               for group in consistency['groups'])
    summary = json.loads((track / 'history/golden_judge_fix_summary.json').read_text())
    archive = track / 'history' / summary['history']
    assert (archive / 'golden/reward.json').is_file()
    assert len(list((archive / 'judge_runs').glob('golden_r[1-5]/reward-details.json'))) == 5
    for folder in sorted((track / 'rollouts').iterdir()):
        candidate = folder / 'candidate.xlsx'
        result = engine.score(engine.extract(candidate, key), key)
        receipt = json.loads((folder / 'reward.json').read_text())
        saved_program = receipt.get('programmatic_score', receipt['reward'])
        assert abs(result['score'] - saved_program) < 0.000001, folder.name
        assert receipt['reward'] == summary['new_model_scores'][folder.name], folder.name
    print(json.dumps({'saved_report_checked': track.parent.name, 'programmatic_golden': 1.0,
                      'valid_new_judgments': 0, 'current_golden_score': None,
                      'archived_judgments_excluded': 5}))


def import_module(name, path):
    specification = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(module)
    return module


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--trajectory', type=Path)
    arguments = parser.parse_args()
    task = Path(__file__).resolve().parents[1]
    track = arguments.trajectory or task.parents[1] / 'ZCloud-财务分析-002-20261020-轨迹报告/cno-insurance-benefits-q1-2025'
    counts = check_saved_cases(task)
    with tempfile.TemporaryDirectory() as temporary:
        scoring_main(temporary)
        saved_main(track)
        access_main()
    check_saved_report(task, track)
    print(json.dumps(dict(status='PASS', **counts)))


if __name__ == '__main__':
    main()
