"""Offline interface and instruction-constraint regression workbooks."""
import argparse
import copy
import importlib.util
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.dont_write_bytecode = True
TASK = Path(__file__).resolve().parents[1]


def load_module(name, path):
    specification = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(module)
    return module


def reverse_columns(worksheet):
    values = [[copy.copy(cell) for cell in row] for row in worksheet]
    for row_number, row in enumerate(values, 1):
        for column_number, cell in enumerate(reversed(row), 1):
            destination = worksheet.cell(row_number, column_number)
            destination.value = cell.value
            destination.number_format = cell.number_format


def swap_rows(worksheet):
    rows = [[copy.copy(cell) for cell in worksheet[row]] for row in (2, 3)]
    for row_number, row in zip((3, 2), rows):
        for column_number, cell in enumerate(row, 1):
            worksheet.cell(row_number, column_number, cell.value)
            worksheet.cell(row_number, column_number).number_format = cell.number_format


def mutate(workbook, case):
    worksheet = workbook[case.get("sheet", "Results")]
    mutation = case["mutation"]
    if mutation == "review_wrong_names_reverse_all":
        workbook["Results"].title = "WrongSheet1"
        workbook["Analysis"].title = "WrongSheet2"
        for current in workbook:
            reverse_columns(current)
    elif mutation == "rename":
        worksheet.title = "Wrong_" + worksheet.title
    elif mutation == "reverse_columns":
        reverse_columns(worksheet)
    elif mutation == "extra_header":
        worksheet.cell(1, worksheet.max_column + 1, "extra")
    elif mutation == "missing_header":
        worksheet.cell(1, worksheet.max_column).value = None
    elif mutation == "shift_header":
        worksheet.insert_rows(1)
    elif mutation == "header_case":
        worksheet["A1"] = worksheet["A1"].value.upper()
    elif mutation == "set_cell":
        worksheet[case["cell"]] = case["value"]
        if "format" in case:
            worksheet[case["cell"]].number_format = case["format"]
    elif mutation == "swap_rows":
        swap_rows(worksheet)
    elif mutation in {"extra_record", "duplicate_record", "blank_key_record"}:
        values = [cell.value for cell in worksheet[2]]
        if mutation == "extra_record":
            values[0] = "Unexpected business record"
        if mutation == "blank_key_record":
            values[0] = None
        worksheet.append(values)
    elif mutation == "missing_record":
        worksheet.delete_rows(2)
    else:
        raise ValueError(mutation)


def definitions():
    cases = [{"id": "review_wrong_structure", "mutation": "review_wrong_names_reverse_all",
              "gate": True, "old_score_from_review": 1.0}]
    for sheet in ("Results", "Analysis", "Summary"):
        for mutation in ("rename", "reverse_columns", "missing_header", "shift_header", "header_case"):
            cases.append({"id": sheet.lower() + "_" + mutation, "sheet": sheet,
                          "mutation": mutation, "gate": mutation == "rename",
                          "cap": mutation != "rename"})
        for mutation in ("swap_rows", "extra_record", "duplicate_record", "blank_key_record", "missing_record"):
            cases.append({"id": sheet.lower() + "_" + mutation, "sheet": sheet,
                          "mutation": mutation, "cap": mutation in {"extra_record", "blank_key_record"},
                          "penalty": "O01" if mutation in {"swap_rows", "missing_record"} else "X2"})
    for label, cell, value, penalty in [
        ("results_empty_source_d2", "D2", None, "S01"),
        ("results_empty_source_d3", "D3", " ", "S01"),
        ("results_numeric_source", "D2", 123, "S01"),
        ("results_generic_source", "D2", "CNO filings", "S02"),
        ("results_fake_file", "D2", "Imaginary.xlsx T010!G19", "S02"),
        ("results_fake_table", "D2", "CNO_2025-03-31_10Q.xlsx T999!G19", "S02"),
        ("results_fake_cell", "D2", "CNO_2025-03-31_10Q.xlsx T010!ZZ99999", "S02"),
        ("results_wrong_unit", "C2", "USD thousands", "U01"),
        ("results_unit_case", "C3", "Percent", "U01"),
        ("results_percentage_general", "B3", 12.13850088530396, "DP2"),
        ("results_percentage_unrounded_coarse_format", "B3", 12.13850088530396, "DP2"),
        ("results_amount_extra_decimal", "B2", 61.71, "DP1"),
        ("results_percentage_outside_tolerance", "B3", 12.1441, "R02"),
    ]:
        case = {"id": label, "mutation": "set_cell", "cell": cell, "value": value, "penalty": penalty,
                "cap": penalty in {"S01", "U01"}}
        if cell.startswith("B"):
            case["format"] = "0.0" if "coarse_format" in label else "0.00" if "tolerance" in label else "General"
        cases.append(case)
    for label, cell, value, penalty in [
        ("analysis_fake_file", "E2", "Imaginary.xlsx", "SD1"),
        ("analysis_empty_location", "F3", None, "SP1"),
        ("analysis_fake_table", "F2", "T999", "SP1"),
        ("analysis_wrong_unit", "D2", "USD thousand", "U01"),
    ]:
        cases.append({"id": label, "sheet": "Analysis", "mutation": "set_cell", "cell": cell,
                      "value": value, "penalty": penalty, "cap": label in {"analysis_empty_location", "analysis_wrong_unit"}})
    for cell in ("E2", "E3", "F2", "F3"):
        for label, value in (("empty", None), ("numeric", 123), ("numeric_text", "123")):
            cases.append({"id": "analysis_" + cell.lower() + "_" + label, "sheet": "Analysis",
                          "mutation": "set_cell", "cell": cell, "value": value, "penalty": "S01", "cap": True})
    return cases


def verify_finalize(finalizer, result, weights):
    criteria = [{"name": name, "weight": abs(weights[name]),
                 "value": 1 - verdict["credit"] if weights[name] < 0 else verdict["credit"],
                 "negate": weights[name] < 0} for name, verdict in result["verdicts"].items()]
    details = {"reward": [{"kind": "programmatic", "criteria": criteria}],
               "interface_gate": result["interface_gate"], "post_pool_rules": result["post_pool_rules"]}
    with tempfile.TemporaryDirectory() as temporary:
        graded = Path(temporary) / "reward.json"
        graded.write_text(json.dumps({"programmatic": result["score"]}))
        graded.with_name("reward-details.json").write_text(json.dumps(details))
        verdict = finalizer.compute(argparse.Namespace(graded=str(graded), graded_rc=0))
        assert verdict["reward"] == result["score"] and verdict["verifier_error"] == 0
        details["reward"].append({"kind": "agent", "criteria": [
            {"id": "L01", "value": 1.0, "raw": "no", "weight": 3, "negate": True,
             "reasoning": "无法解析候选文件，文件无法解析，无法检查来源。"}]})
        graded.with_name("reward-details.json").write_text(json.dumps(details))
        unavailable = finalizer.compute(argparse.Namespace(graded=str(graded), graded_rc=0))
        assert unavailable["reward"] == 0 and unavailable["verifier_error"] == 1
        assert unavailable["llm_evidence_errors"]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--baseline", action="store_true")
    arguments = parser.parse_args()
    engine = load_module("regression_engine", TASK / "tests/score_output.py")
    finalizer = load_module("regression_finalizer", TASK / "tests/finalize.py")
    key = json.loads((TASK / "tests/results_key.json").read_text())
    weights = {item["criterion_id"]: item["weight"] for item in key["rows"] + key["tables"] + key["checks"]}
    directory = TASK / "validation/negative_cases"
    directory.mkdir(parents=True, exist_ok=True)
    metadata = directory / "cases.json"
    previous_metadata = json.loads(metadata.read_text()) if metadata.exists() else {}
    previous = {case["id"]: case for case in previous_metadata.get("cases", []) + previous_metadata.get("controls", [])}
    golden = TASK / "solution/golden/answer.xlsx"
    baseline = engine.score(engine.extract(golden, key), key)
    assert baseline["score"] == 1.0, baseline
    results = []
    for definition in definitions():
        workbook = load_workbook(golden)
        mutate(workbook, definition)
        path = directory / (definition["id"] + ".xlsx")
        workbook.save(path)
        workbook.close()
        result = engine.score(engine.extract(path, key), key)
        verify_finalize(finalizer, result, weights)
        pooled, counted, broken = finalizer.pooled_score({"reward": {"criteria": [
            {"weight": abs(weights[name]), "value": 1 - verdict["credit"] if weights[name] < 0 else verdict["credit"],
             "negate": weights[name] < 0} for name, verdict in result["verdicts"].items()]}})
        pooled = finalizer.apply_post_pool_rules(result, pooled)
        assert not broken and counted == len(weights) and abs(pooled - result["score"]) < 0.000001
        failed = [name for name, verdict in result["verdicts"].items()
                  if verdict["credit"] == (1 if weights[name] < 0 else 0)]
        record = dict(definition, workbook=path.name,
                      expected_outcome="gate 0" if definition.get("gate") else "cap <=0.69" if definition.get("cap") else "penalty <1",
                      old_score_from_review=definition.get("old_score_from_review"),
                      old_program_score=previous.get(definition["id"], {}).get("old_program_score"),
                      round1_program_score=previous.get(definition["id"], {}).get("round1_program_score", previous.get(definition["id"], {}).get("new_program_score")),
                      new_program_score=None if arguments.baseline else result["score"],
                      triggered=failed,
                      interface_gate=result.get("interface_gate"), post_pool_rules=result.get("post_pool_rules"))
        if not arguments.baseline:
            assert result["score"] < 1.0, record
            if definition.get("gate"):
                assert result["score"] == 0 and result["interface_gate"]["failed"], record
            elif definition.get("cap"):
                assert result["score"] <= 0.69 and any(rule["triggered"] for rule in result["post_pool_rules"]), record
            if definition.get("penalty"):
                assert definition["penalty"] in failed, record
        results.append(record)
    controls = []
    if not arguments.baseline:
        for label in ("support_sheet_unconstrained", "rounded_stored_values", "percentage_excel_format", "styling_only_extra_column", "trailing_header_spaces", "leading_header_spaces", "results_percentage_wrong_format", "formatted_unrounded_value", "results_extra_header", "analysis_extra_header", "summary_extra_header", "trailing_duplicate_header"):
            workbook = load_workbook(golden)
            if label == "support_sheet_unconstrained":
                support = workbook.create_sheet("Support")
                support.append(["Anything", "#DIV/0!", 123.4567])
            elif label == "rounded_stored_values":
                workbook["Results"]["B2"].number_format = "General"
                workbook["Results"]["B3"] = 12.14
                workbook["Results"]["B3"].number_format = "General"
            elif label == "percentage_excel_format":
                workbook["Results"]["B3"] = 61.7 / 508.3
                workbook["Results"]["B3"].number_format = "0.00%"
            elif label == "styling_only_extra_column":
                workbook["Results"]["E1"].number_format = "0.00"
            elif label in {"trailing_header_spaces", "leading_header_spaces"}:
                for worksheet in workbook:
                    for cell in worksheet[1]:
                        cell.value = cell.value + "  " if label == "trailing_header_spaces" else "  " + cell.value
            elif label == "results_percentage_wrong_format":
                workbook["Results"]["B3"] = 12.14
                workbook["Results"]["B3"].number_format = "0.0"
            elif label == "formatted_unrounded_value":
                workbook["Results"]["B2"] = 61.71
                workbook["Results"]["B2"].number_format = "0.0"
            elif label == "trailing_duplicate_header":
                for worksheet, header in [(workbook["Results"], "value"), (workbook["Analysis"], "amount"), (workbook["Summary"], "insurance_policy_benefits")]:
                    extra_column = worksheet.max_column + 1
                    worksheet.cell(1, extra_column, header)
                    worksheet.cell(2, extra_column, 123)
                    worksheet.cell(3, extra_column, 456)
            else:
                worksheet = workbook[label.split("_")[0].title()]
                worksheet.cell(1, worksheet.max_column + 1, "extra")
                worksheet.cell(2, worksheet.max_column, "Additional permitted support")
            path = directory / (label + ".xlsx")
            workbook.save(path)
            workbook.close()
            score = engine.score(engine.extract(path, key), key)["score"]
            assert score == 1.0, (label, score)
            controls.append({"id": label, "workbook": path.name, "expected_program_score": 1.0, "program_score": score,
                             "round1_program_score": previous.get(label, {}).get("round1_program_score", previous.get(label, {}).get("new_program_score"))})
    metadata.write_text(json.dumps({"scope": "deterministic programmatic subset; no new text judgment",
                                    "precision_rule": "p decimals satisfied by number_format OR stored value equal to rounding to p decimals; percentage formats use percentage points for stored comparison",
                                    "old_score_scope": "review value only where present; other old_program_score values are offline v1 replays on the formatted golden",
                                    "golden_program_score": baseline["score"], "cases": results, "controls": controls}, ensure_ascii=False, indent=2) + "\n")
    verify_finalize(finalizer, baseline, weights)
    print(json.dumps({"negative_cases": len(results), "controls": len(controls), "golden": baseline["score"],
                      "finalize_policy_and_unreadable_judge_checks": len(results) + 1,
                      "mode": "baseline" if arguments.baseline else "verified"}))


if __name__ == "__main__":
    main()
